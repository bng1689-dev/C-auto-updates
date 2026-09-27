#!/usr/bin/env python3
"""ถ่ายภาพหน้าจอแถบสถานะในทุกสถานะ — ไว้ตรวจงานออกแบบด้วยตา (ไม่ใช่ชุดทดสอบ)

    python tools/shots.py            → build/shots/*.png

ใช้แอปจริงจาก build/ (เหมือนชุดทดสอบ) แต่แทนสถานะรอบทำงานด้วยค่าปลอม จึงดูได้ทุกสถานะ
โดยไม่ต้องเปิด Chrome ค้นจริง · ต้องมี Chromium ของ Playwright (python -m playwright install chromium)
"""
import logging
import os
import shutil
import sys
import tempfile
import threading
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_shot_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA)
sys.path.insert(0, str(ROOT / "tests"))
from _app import CHROME  # noqa: E402
from backend import server  # noqa: E402
import engine  # noqa: E402

OUT = ROOT / "build" / "shots"
PORT = 8809
ST = {}


def fake_status():
    base = {"state": "idle", "browser_open": False, "percent": 0, "processed": 0, "total_rows": 0,
            "message": "", "incomplete": 0, "no_charge": 0, "retry_pending": 0, "retry_pass": 0,
            "errors": 0, "found": 0, "notfound": 0, "restored": 0, "awaiting_step": False,
            "step_mode": False, "current_step": "", "speed": 3, "speed_label": "ปกติ",
            "log": ["เปิดเบราว์เซอร์แล้ว"], "save_seq": 0, "file_index": 1, "file_total": 1,
            "current_file": "", "current_path": "", "row": 0, "file_done": 0, "errors_left": 0,
            "account": "", "recording": False, "recording_steps": 0, "pause_reason": "",
            "autostart_in": 0, "login_detected": False, "padded": 0, "invalid_ids": 0}
    base.update(ST)
    return base


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    logging.getLogger("werkzeug").setLevel(logging.ERROR)
    server.app.config["TESTING"] = False
    server.manager.status = fake_status
    server.manager.current_user_id = 1
    threading.Thread(target=lambda: server.app.run(port=PORT, threaded=True, use_reloader=False),
                     daemon=True).start()
    time.sleep(1.2)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROME, headless=True, args=["--disable-background-networking"])
        page = b.new_page(viewport={"width": 1280, "height": 900})
        errs = []
        page.on("pageerror", lambda e: errs.append(str(e)))
        page.goto(f"http://127.0.0.1:{PORT}", wait_until="domcontentloaded")
        page.wait_for_selector("#authSetupBox:not(.hidden)", timeout=8000)
        page.fill("#suUser", "admin1")
        page.fill("#suName", "แอดมิน")
        page.fill("#suPw", "secret9")
        page.fill("#suPw2", "secret9")
        page.click("#authSetupForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(1, 2, "เลขบัตร")
        for i in range(120):
            ws.cell(2 + i, 2, f"1{str(i).zfill(12)}")
            if i < 57:
                ws.cell(2 + i, 6, " - ")
        p = TEST_DATA / "รายชื่อผู้สมัคร_กันยายน.xlsx"
        wb.save(p)
        engine.save_queue([{"id": "q1", "path": str(p), "user_id": 1}])
        page.click('[data-view="run"]')
        page.wait_for_selector("#queueList .qitem", timeout=6000)

        def shot(name, wait=1600):
            page.wait_for_timeout(wait)
            page.locator("#runBanner").screenshot(path=str(OUT / f"{name}.png"))
            print("  บันทึก", OUT / f"{name}.png")

        ST.update(state="awaiting_login", browser_open=True, autostart_in=7, login_detected=True,
                  message="เข้าสู่ระบบแล้ว — เริ่มค้นอัตโนมัติใน 7 วินาที (กด 'เริ่มค้นหา' เพื่อเริ่มทันที)")
        page.wait_for_selector("#runBanner:not(.hidden)", timeout=8000)
        shot("1_awaiting_login")
        ST.update(state="running", autostart_in=0, percent=52, file_done=62, total_rows=120, row=63,
                  found=18, notfound=41, errors=3, restored=5, incomplete=2, no_charge=1, retry_pending=3,
                  current_file=p.name, current_path=str(p),
                  message="กำลังค้นหา — ดูรายละเอียดแต่ละแถวในบันทึกการทำงานด้านล่าง")
        shot("2_running", 2200)
        ST.update(file_done=66, percent=55, row=67, found=20)
        page.wait_for_timeout(5500)
        ST.update(file_done=70, percent=58, row=71, found=22)
        shot("3_running_eta", 1800)
        ST.update(state="paused", pause_reason="manual",
                  message="⏸ พักชั่วคราวตามคำขอ · ค้างที่แถว 71 — ผลที่ค้นแล้วถูกบันทึกไว้ / กด 'ค้นต่อ ▶' เมื่อพร้อม")
        shot("4_paused")
        ST.update(state="done", percent=100, file_done=120, row=120, found=40, notfound=77, errors=3, errors_left=3,
                  message="เสร็จสมบูรณ์ — ค้น 120, พบ 40, ไม่พบ 77, ผิดพลาด 3 · ⚠ ยังมีแถวผิดพลาดค้าง 3 แถว")
        shot("5_done")
        ST.update(state="error", browser_open=False,
                  message="หน้าต่าง Chrome ถูกปิด — ผลที่ค้นแล้วถูกบันทึกไว้ กด ▶ LOGIN CRIMES เพื่อค้นต่อจากแถวเดิม")
        shot("6_error")
        engine.save_queue([])          # ใบเดียวเต็มจอ (ไม่มีคิว)
        ST.update(state="running", browser_open=True, percent=58, errors_left=0,
                  message="กำลังค้นหา — ดูรายละเอียดแต่ละแถวในบันทึกการทำงานด้านล่าง")
        page.reload(wait_until="domcontentloaded")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)
        page.click('[data-view="run"]')
        page.wait_for_selector("#runBanner:not(.hidden)", timeout=8000)
        shot("7_single_running", 2500)
        page.screenshot(path=str(OUT / "8_page_running.png"))
        b.close()
        real = [e for e in errs if "favicon" not in e]
        print("JavaScript errors:", real or "ไม่มี")
        return 1 if real else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
