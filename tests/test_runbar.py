#!/usr/bin/env python3
"""ทดสอบแถบสถานะ/ปุ่มควบคุมรอบค้น บนหน้าจอจริง (Chromium)

ผู้ใช้รายงานว่า "ไม่มีปุ่มกดเริ่มขึ้นมาให้กด" — ตรวจว่าแต่ละสถานะของรอบทำงาน
ปุ่มที่ควรโผล่ โผล่จริง มองเห็นจริง และกดได้จริง (ไม่ถูกบัง/ไม่หลุดออกนอกจอ)
โดยเฉพาะหลัง v3.2.0 ที่ย้ายแถบสถานะเข้าไปอยู่ในแถวเดียวกับกราฟคิว (.run-row)
"""
import os
import shutil
import sys
import tempfile
import threading
import time
from datetime import datetime
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_runbar_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA)
from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))
from backend import db, engine, server  # noqa: E402

PORT = 8797
PASS = FAIL = 0
JS_ERRORS = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


STATE = {"state": "idle"}


def fake_status():
    s = {"state": STATE["state"], "browser_open": STATE["state"] not in ("idle",),
         "percent": 40, "processed": 4, "total_rows": 10, "message": "ทดสอบ",
         "incomplete": 0, "retry_pending": 0, "retry_pass": 0, "errors": 0,
         "awaiting_step": STATE.get("awaiting_step", False),
         "step_mode": STATE.get("step_mode", False),
         "current_step": "กรอกเลขบัตร", "speed": 3, "speed_label": "ปกติ",
         "log": ["เปิดเบราว์เซอร์แล้ว"], "save_seq": 0, "file_index": 1,
         "file_done": STATE.get("file_done", 0), "current_file": "งาน.xlsx", "current_path": "งาน.xlsx",
         "account": "", "recording": False, "recording_steps": 0, "pause_reason": ""}
    return s


def vis(page, sel):
    """มองเห็นจริงไหม: ไม่ถูกซ่อน มีขนาด และจุดกึ่งกลางเป็นตัวมันเอง (ไม่มีอะไรบัง)"""
    return page.evaluate(
        """(sel) => {
          const el = document.querySelector(sel);
          if (!el) return 'ไม่มีองค์ประกอบนี้';
          if (el.classList.contains('hidden')) return 'ถูกซ่อน (class hidden)';
          const r = el.getBoundingClientRect();
          if (r.width < 1 || r.height < 1) return 'ขนาดเป็นศูนย์';
          if (r.bottom < 0 || r.top > innerHeight * 4) return 'อยู่นอกหน้าจอมาก';
          const top = document.elementFromPoint(r.left + r.width / 2, r.top + r.height / 2);
          if (top && !el.contains(top) && !top.contains(el)) return 'ถูกบังด้วย ' + top.tagName;
          return '';
        }""", sel)


def main():
    server.app.config["TESTING"] = False
    server.manager.status = fake_status
    server.manager.current_user_id = 1
    threading.Thread(target=lambda: server.app.run(port=PORT, threaded=True,
                                                   use_reloader=False), daemon=True).start()
    time.sleep(1.2)

    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROME, headless=True,
                               args=["--disable-background-networking"])
        page = b.new_page(viewport={"width": 1280, "height": 900})
        page.on("pageerror", lambda e: JS_ERRORS.append(f"pageerror: {e}"))
        page.goto(f"http://127.0.0.1:{PORT}", wait_until="domcontentloaded")
        page.wait_for_selector("#authSetupBox:not(.hidden)", timeout=8000)
        page.fill("#suUser", "admin1")
        page.fill("#suName", "แอดมิน")
        page.fill("#suPw", "secret9")
        page.fill("#suPw2", "secret9")
        page.click("#authSetupForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)

        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.cell(1, 2, "เลขบัตร")
        for i in range(4):
            ws.cell(2 + i, 2, f"1{str(i).zfill(12)}")
        p = TEST_DATA / "งาน.xlsx"; wb.save(p)
        engine.save_queue([{"id": "q1", "path": str(p), "user_id": 1}])

        page.click('[data-view="run"]')
        page.wait_for_selector("#view-run:not(.hidden)", timeout=5000)
        page.wait_for_selector("#queueList .qitem", timeout=6000)

        # ── สถานะ: รอยืนยัน login → ต้องเห็นปุ่ม 'เริ่มค้นหา' ──
        STATE["state"] = "awaiting_login"
        page.wait_for_selector("#runBanner:not(.hidden)", timeout=8000)
        page.wait_for_timeout(500)
        check("แถบสถานะโผล่เมื่อรอยืนยัน login", vis(page, "#runBanner") == "",
              vis(page, "#runBanner"))
        why = vis(page, "#btnConfirmLogin")
        check("ปุ่ม 'เริ่มค้นหา' โผล่และมองเห็นจริง", why == "", why)
        check("แถวคู่ไม่ถูกซ่อนทั้งแถว",
              page.eval_on_selector(".run-row", "el => !el.classList.contains('is-empty')"))
        check("กราฟคิวและแถบสถานะอยู่แถวเดียวกัน คนละครึ่ง", page.evaluate(
            """() => { const a = document.querySelector('#runStatusCard').getBoundingClientRect();
                 const b = document.querySelector('#runBanner').getBoundingClientRect();
                 return Math.abs(a.top - b.top) < 40 && a.right <= b.left + 2
                        && Math.abs(a.width - b.width) < 40; }"""))
        page.click("#btnConfirmLogin")      # ต้องกดได้จริง ไม่โดนบัง
        check("กดปุ่ม 'เริ่มค้นหา' ได้จริง", True)

        # ── สถานะอื่น ๆ ──
        STATE["state"] = "running"
        page.wait_for_timeout(1600)
        check("กำลังค้น: ปุ่ม 'เริ่มค้นหา' หายไป",
              page.eval_on_selector("#btnConfirmLogin", "el => el.classList.contains('hidden')"))
        for sel, label in (("#btnPause", "พัก"), ("#btnStop", "หยุด"),
                           ("#btnSaveNow", "บันทึก·ดาวน์โหลด")):
            why = vis(page, sel)
            check(f"กำลังค้น: ปุ่ม '{label}' มองเห็นจริง", why == "", why)

        STATE.update(state="running", step_mode=True, awaiting_step=True)
        page.wait_for_timeout(1600)
        why = vis(page, "#btnNextStep")
        check("โหมดทีละขั้น: ปุ่ม 'ทำขั้นต่อไป' มองเห็นจริง", why == "", why)

        STATE.update(state="paused", awaiting_step=False)
        page.wait_for_timeout(1600)
        why = vis(page, "#btnResume")
        check("พักอยู่: ปุ่ม 'ค้นต่อ' มองเห็นจริง", why == "", why)

        # ── v3.6.1: เวลาที่เหลือต้องเริ่มนับใหม่หลังพัก — ช่วงที่พักไม่ถูกนับเป็นเวลาค้น ──
        # เดิมจุดตั้งต้นค้างไว้ข้ามช่วงพัก → กดค้นต่อแล้วอัตราแถว/นาทีต่ำเกินจริง เวลาที่เหลือพุ่งเกินจริง
        STATE.update(state="running", step_mode=False, awaiting_step=False, file_done=10)
        page.wait_for_timeout(1600)                       # จุดตั้งต้นของอัตรา = 10 แถว
        STATE.update(file_done=14)
        page.wait_for_timeout(5500)                       # +4 แถวใน >5 วิ → มีค่า ETA ให้โชว์
        eta = page.eval_on_selector("#rbEta", "el => el.classList.contains('hidden') ? '' : el.textContent")
        check("กำลังค้น: ชิปเวลาที่เหลือโผล่พร้อมอัตราแถว/นาที", "แถว/นาที" in eta, eta)
        STATE.update(state="paused")
        page.wait_for_timeout(1600)
        STATE.update(state="running")                     # ค้นต่อ — ยังไม่มีแถวใหม่
        page.wait_for_timeout(1600)
        check("ค้นต่อหลังพัก: ชิปเวลาที่เหลือซ่อนจนกว่าจะมีอัตราใหม่ (ไม่เอาช่วงพักมาคิด)",
              page.eval_on_selector("#rbEta", "el => el.classList.contains('hidden')"))
        STATE.update(file_done=17)
        page.wait_for_timeout(5500)
        eta = page.eval_on_selector("#rbEta", "el => el.classList.contains('hidden') ? '' : el.textContent")
        check("ค้นต่อไปอีก 3 แถว → เวลาที่เหลือกลับมาโชว์จากอัตราใหม่", "แถว/นาที" in eta, eta)

        # ── ไม่มีคิวเลย + ไม่มีรอบทำงาน → แถวคู่ต้องซ่อนหมด ไม่เหลือช่องว่าง ──
        STATE.update(state="idle", step_mode=False)
        engine.save_queue([])
        page.reload(wait_until="domcontentloaded")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)
        page.click('[data-view="run"]')
        page.wait_for_timeout(1800)
        check("ว่างทั้งคู่ → ซ่อนทั้งแถว ไม่เหลือช่องว่าง",
              page.eval_on_selector(".run-row", "el => el.classList.contains('is-empty')"))

        # ── v3.5.0: session หมดอายุกลางคัน (PASSCODE ครบ 12 ชม.) แล้วเข้าสู่ระบบใหม่โดยไม่รีโหลดหน้า ──
        # เดิม poll หยุดตอนได้ 401 แล้วไม่กลับมาอีก → กด LOGIN CRIMES แล้วแถบสถานะ/ปุ่มเริ่มไม่โผล่เลย
        engine.save_queue([{"id": "q1", "path": str(p), "user_id": 1}])
        page.evaluate("() => fetch('/api/logout', {method: 'POST'})")      # ล้าง session ฝั่งเซิร์ฟเวอร์
        page.wait_for_selector("#authLoginBox:not(.hidden)", timeout=8000)
        check("session หมดอายุ → กลับไปหน้าเข้าสู่ระบบ", True)
        page.fill("#authLoginUser", "admin1")
        page.fill("#authLoginPw", "secret9")
        page.click("#authLoginForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)
        page.click('[data-view="run"]')
        STATE.update(state="awaiting_login", step_mode=False, awaiting_step=False)
        try:
            page.wait_for_selector("#btnConfirmLogin:not(.hidden)", timeout=6000)
            ok = vis(page, "#btnConfirmLogin") == ""
        except Exception:
            ok = False
        check("เข้าสู่ระบบใหม่แล้ว poll กลับมาทำงาน → ปุ่ม 'เริ่มค้นหา' โผล่", ok)
        hits = []
        page.on("request", lambda r: hits.append(1) if r.url.endswith("/api/status") else None)
        page.wait_for_timeout(6000)
        check("poll ไม่ซ้อนสองวง (ราว 1 ครั้ง/1.2 วิ)", 3 <= len(hits) <= 7, f"{len(hits)} ครั้งใน 6 วิ")

        b.close()

    real = [e for e in JS_ERRORS if "favicon" not in e]
    check("ไม่มี JavaScript error", not real, "\n    ".join(real[:5]))
    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
