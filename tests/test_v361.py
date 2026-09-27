#!/usr/bin/env python3
"""v3.6.1 — เลขบัตรอยู่คอลัมน์ F (ช่องผล) ต้องถูกปฏิเสธ ไม่ใช่ขึ้น 'เสร็จ' โดยไม่ได้ค้น (ไม่เปิดเบราว์เซอร์)

ข้อค้นพบจากรีวิว v3.6.0: เมื่อผลอยู่ F เสมอ ไฟล์ที่เลขบัตรอยู่ F จะทำให้ทุกแถวถูกนับว่า 'มีผลแล้ว'
(เลขบัตรอยู่ในช่องผล) → รอบค้นข้ามทุกแถว · หน้าคิวขึ้น File ready · ถูกตัดออกจากคิว ทั้งที่ไม่ได้ค้นสักแถว
  • engine.column_clash บอกว่าคอลัมน์ชนกัน
  • /api/upload ปฏิเสธ (ระบุเอง / ตรวจจับได้ F) พร้อมข้อความบอกวิธีแก้ และลบไฟล์ที่เพิ่งอัปโหลด
  • รูปแบบที่จำไว้จากรุ่นก่อนว่าเลขบัตรอยู่ F → ไม่ใช้ ตรวจจับใหม่แทน
  • งานเก่าในคิวที่ id_col=F: /api/queue รายงาน invalid ไม่ใช่ done · prune ไม่ตัดออก · รอบค้นข้ามพร้อมบันทึก
  • หน้าจอ: ป้าย 'ต้องแก้ไฟล์' และ ETA ล้างจุดตั้งต้นตอนพัก
"""
import io
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v361_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "data")
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "up")
from _app import APP  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)

from backend import server  # noqa: E402
import db  # noqa: E402
import engine  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def make_xlsx(name, headers, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(headers, 1):
        if h is not None:
            ws.cell(1, c, h)
    for r, vals in enumerate(rows, 2):
        for c, v in enumerate(vals, 1):
            if v is not None:
                ws.cell(r, c, v)
    p = TEST_DATA / name
    wb.save(p)
    return p


def upload(c, path, **form):
    data = {"file": (io.BytesIO(path.read_bytes()), path.name)}
    data.update(form)
    r = c.post("/api/upload", data=data, content_type="multipart/form-data")
    return r.status_code, (r.get_json() or {})


def main():
    print("── ตัวตรวจคอลัมน์ชนกัน ──")
    check("F ชน", bool(engine.column_clash("F")) and bool(engine.column_clash(" f ")))
    check("B / ว่าง ไม่ชน", engine.column_clash("B") is None and engine.column_clash("") is None)
    check("ข้อความบอกวิธีแก้", "คอลัมน์เขียนผล" in engine.column_clash("F") and "อัปโหลดใหม่" in engine.column_clash("F"))

    print("\n── อัปโหลด ──")
    server.app.config["TESTING"] = True
    c = server.app.test_client()
    c.post("/api/setup", json={"username": "admin1", "password": "secret9",
                               "password2": "secret9", "display_name": "Admin"})
    up_dir = Path(os.environ["CRIMES_UPLOAD_DIR"])
    pa = make_xlsx("a.xlsx", ["ลำดับ", "เลขบัตร"], [[1, 3101501909804], [2, 3101501909805]])
    pf = make_xlsx("f.xlsx", ["ลำดับ", None, None, None, None, "เลขบัตร"],
                   [[1, None, None, None, None, 3101501909804], [2, None, None, None, None, 3101501909805]])
    engine.save_queue([])
    code, j = upload(c, pa, id_col="F")
    check("ระบุเองว่าเลขบัตรอยู่ F → ปฏิเสธพร้อมบอกวิธีแก้", code == 400 and "คอลัมน์เขียนผล" in j.get("error", ""),
          f"{code} {j}")
    check("ไฟล์ที่ถูกปฏิเสธไม่ค้างในโฟลเดอร์อัปโหลด", not list(up_dir.rglob("a.xlsx")))
    code, j = upload(c, pf)
    check("ตรวจจับได้ว่าเลขบัตรอยู่ F → ปฏิเสธ", code == 400 and "คอลัมน์เขียนผล" in j.get("error", ""), f"{code} {j}")
    check("ไฟล์ที่ถูกปฏิเสธไม่ค้างในโฟลเดอร์อัปโหลด (2)", not list(up_dir.rglob("f.xlsx")))
    check("คิวยังว่าง ไม่มีงานที่ค้นไม่ได้หลุดเข้ามา", engine.load_queue() == [])
    # รูปแบบที่จำไว้จากรุ่นก่อนบอกว่าเลขบัตรอยู่ F แต่ไฟล์จริงเลขอยู่ B → ต้องตรวจจับใหม่ ไม่ปฏิเสธ
    sig, _ = engine.read_header_signature(pa)
    db.save_file_format(sig, "F", "F", 2, "")
    code, j = upload(c, pa)
    check("รูปแบบที่จำไว้บอก F → ไม่ใช้ ตรวจจับใหม่ได้ B", code == 200 and j.get("job", {}).get("id_col") == "B",
          f"{code} {j}")
    check("จำรูปแบบใหม่เป็น B", (db.get_file_format(sig) or {}).get("id_col") == "B")
    check("ระบุเอง B ทับที่จำไว้ ใช้ได้ตามปกติ", upload(c, pa, id_col="B")[0] == 200)

    print("\n── งานเก่าในคิวที่เลขบัตรอยู่ F ──")
    engine.save_queue([{"id": "old", "path": str(pf), "user_id": 1, "id_col": "F"},
                       {"id": "ok", "path": str(pa), "user_id": 1, "id_col": "B"}])
    q = c.get("/api/queue").get_json()
    old = next(x for x in q if x["id"] == "old")
    ok_ = next(x for x in q if x["id"] == "ok")
    check("/api/queue: งาน F รายงาน invalid ไม่ใช่ done/searched",
          "คอลัมน์เขียนผล" in old.get("invalid", "") and not old["done"] and old["searched"] == 0, str(old))
    check("/api/queue: งานปกติไม่มี invalid", ok_.get("invalid") == "" and ok_["pending"] == 2, str(ok_))
    r = c.post("/api/queue/prune").get_json()
    check("prune ไม่ตัดงาน F ออก (ผู้ใช้ต้องเห็นว่าต้องแก้ไฟล์)", r["removed"] == 0 and len(engine.load_queue()) == 2, str(r))
    n0 = len(server.manager.log_lines)
    server.manager._process_file(None, {"id": "old", "path": str(pf), "user_id": 1, "id_col": "F"})
    new_log = server.manager.log_lines[n0:]
    check("รอบค้นข้ามไฟล์ F ก่อนแตะเบราว์เซอร์ พร้อมบันทึกเหตุผล",
          any("ข้าม" in ln and "คอลัมน์เขียนผล" in ln for ln in new_log) and server.manager.current_file == "",
          str(new_log))
    pr = engine.scan_progress(pf, "F", engine.RESULT_COL, 2)
    check("(ยืนยันเหตุผลของกติกา) อ่านตรง ๆ ทุกแถวจะนับว่ามีผลแล้ว", pr["ok"] == 2 and pr["pending"] == 0, str(pr))

    print("\n── หน้าจอ ──")
    html = (APP / "frontend" / "index.html").read_text(encoding="utf-8")
    check("หน้าคิวมีป้าย 'ต้องแก้ไฟล์' จาก invalid", "j.invalid" in html and "ต้องแก้ไฟล์" in html)
    check("ETA ล้างจุดตั้งต้นตอนพัก",
          re.search(r'if\(s\.state === "paused"\)\{[^}]*_eta = \{key:""', html) is not None)

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
