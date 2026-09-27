#!/usr/bin/env python3
"""v3.5.0 — ทดสอบกติกาใหม่ระดับหน่วย (ไม่เปิดเบราว์เซอร์)

  • ตัวอ่านเลขบัตร: เติม 0 เฉพาะ 12 หลัก, ทศนิยม, เลขไทย, ข้อความย่อ E+, ค่าแปลก ๆ
  • จัดกลุ่ม error: 'หน้าเว็บเปลี่ยนไป/ไม่พบช่องกรอก' ต้องลองใหม่ได้ (เดิมถาวร) · หลุด login = auth
  • ความคืบหน้าไฟล์อ่านจากไฟล์จริง (ok/errors/pending/missing/backup)
  • คิวเขียนแบบ atomic + อ่านไฟล์คิวเพี้ยนไม่พัง
  • API: /api/queue มีสถานะใหม่ · /api/queue/prune · start/retry ตอบ error ชัดเจน
  • เพดานลองใหม่ 'ต่อรอบ' — แถว ERROR#3 / ERROR! จากรอบก่อนถูกเก็บมาลองใหม่
"""
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v350_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "data")
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "up")
from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))

from backend import server  # noqa: E402
import engine  # noqa: E402
import worker  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def xlsx(name, ids, results=None, out_col=6):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 2, "เลขบัตร")
    for i, v in enumerate(ids):
        ws.cell(2 + i, 2, v)
        if results and i < len(results) and results[i] is not None:
            ws.cell(2 + i, out_col, results[i])
    p = TEST_DATA / name
    wb.save(p)
    return p


def main():
    print("── ตัวอ่านเลขบัตร ──")
    N = engine.normalize_id
    cases = [
        (3101501909804, ("3101501909804", False, False), "int 13 หลัก"),
        (123456789012, ("0123456789012", True, False), "int 12 หลัก → เติม 0 หน้า"),
        (123456789012.0, ("0123456789012", True, False), "ทศนิยม 12 หลัก → เติม 0 หน้า (ไม่ต่อท้าย)"),
        (3101501909804.0, ("3101501909804", False, False), "ทศนิยม 13 หลัก ใช้ได้ (เดิมกลายเป็น 14 หลัก)"),
        ("123456789012", ("0123456789012", True, False), "ข้อความ 12 หลัก"),
        (" 1-2345-67890-12 ", ("0123456789012", True, False), "มีขีด/ช่องว่าง 12 หลัก"),
        ("3101501909804.0", ("3101501909804", False, False), "ข้อความลงท้าย .0"),
        ("๓๑๐๑๕๐๑๙๐๙๘๐๔", ("3101501909804", False, False), "เลขไทย"),
        ("1.23457E+12", (None, False, True), "ข้อความย่อ E+ (หลักหาย) → ใช้ไม่ได้"),
        (123456789, (None, False, True), "9 หลัก → ใช้ไม่ได้ (เดิมเติม 0 สี่ตัวแล้วค้นมั่ว)"),
        ("12345678901234", (None, False, True), "14 หลัก → ใช้ไม่ได้"),
        (12345.5, (None, False, True), "ทศนิยมไม่ลงตัว → ใช้ไม่ได้"),
        (True, (None, False, True), "ค่า TRUE ในเซลล์ → ใช้ไม่ได้"),
        (None, (None, False, False), "เซลล์ว่าง"),
    ]
    for raw, want, label in cases:
        got = N(raw)
        check(label, got == want, f"{raw!r} → {got}")
    check("หมายเหตุเลขใช้ไม่ได้ไม่ขึ้นต้นด้วย ERROR (จะได้ไม่วนค้นซ้ำ)",
          not engine.invalid_id_note(123).startswith("ERROR")
          and "มี 3 หลัก" in engine.invalid_id_note(123))
    check("ตรวจคอลัมน์นับเลขทศนิยม 13 หลักได้", len(engine._digits_only(3101501909804.0)) == 13)

    print("\n── จัดกลุ่ม error ──")
    C = engine.classify_error
    check("ERROR 'หน้าเว็บเปลี่ยนไป' แบบเดิม → ลองใหม่ได้ (เดิมถาวร)",
          C("หน้าค้นหา CRIMES เปลี่ยนไป — ไม่พบช่องกรอกเลขบัตร (#inputPid) อาจต้องอัปเดต selector") == "retry")
    check("ข้อความใหม่ 'หน้าค้นหายังไม่พร้อม' → ลองใหม่ได้",
          C("หน้าค้นหา CRIMES ยังไม่พร้อม — ไม่พบช่องกรอกเลขบัตร") == "retry")
    check("หลุด login → auth (พักรอคน ไม่เขียน ERROR)",
          C("หลุดการเข้าสู่ระบบ CRIMES (หน้าเว็บกลับไปหน้าเข้าสู่ระบบ) — กรุณาเข้าสู่ระบบใหม่ในหน้าต่าง Chrome") == "auth")
    check("HTTP 401 จากคำตอบการค้น → auth", C("เว็บ CRIMES ตอบการค้นผิดพลาด (HTTP 401)") == "auth")
    check("HTTP 500 → ลองใหม่ได้", C("เว็บ CRIMES ตอบการค้นผิดพลาด (HTTP 500)") == "retry")
    check("เน็ตหลุด → net", C("page.goto: net::ERR_INTERNET_DISCONNECTED") == "net")
    check("ตรวจจับ Chrome ถูกปิด",
          engine.is_browser_closed_error("Target page, context or browser has been closed")
          and not engine.is_browser_closed_error("net::ERR_CONNECTION_CLOSED"))
    info = engine.parse_error_cell("ERROR!: old")
    check("ยังอ่าน ERROR! แบบเก่าได้", info and info["permanent"] is True)

    print("\n── ความคืบหน้าจากไฟล์จริง ──")
    p = xlsx("prog.xlsx", [3101501909800 + i for i in range(5)],
             ["คดีที่1 x ปี2560", " - ", "ERROR#2: t", None, "⚠ เลขบัตรไม่ถูกต้อง (มี 3 หลัก) — ไม่ได้ค้น"])
    pr = engine.scan_progress(p, "B", "F", 2)
    check("นับ ok/errors/pending ถูก", (pr["total"], pr["ok"], pr["errors"], pr["pending"]) == (5, 3, 1, 1),
          str(pr))
    b = p.with_name("prog_backup.xlsx")
    xlsx("prog_backup.xlsx", [3101501909800 + i for i in range(5)], [" - "] * 5)
    os.utime(b, (time.time() + 5, time.time() + 5))
    pr = engine.scan_progress(p, "B", "F", 2)
    check("ไฟล์ต้นฉบับถูกล็อก → อ่านจาก _backup ที่ใหม่กว่า", pr["ok"] == 5 and pr["pending"] == 0, str(pr))
    check("ไฟล์หาย → missing", engine.scan_progress(TEST_DATA / "nope.xlsx")["missing"] is True)
    dup = xlsx("dup.xlsx", [3101501909804] * 3, [" - "] * 3)
    check("เลขซ้ำในไฟล์ไม่ทำให้ค้างเป็นยังไม่ครบ (เดิมนับเลขไม่ซ้ำจากฐานข้อมูล)",
          engine.scan_progress(dup)["ok"] == 3)

    print("\n── คิว ──")
    qp = Path(engine.DEFAULT_QUEUE)
    engine.save_queue([{"id": "a", "path": str(p)}])
    check("เขียนคิวแบบ atomic ไม่มีไฟล์ .tmp ค้าง", not qp.with_name(qp.name + ".tmp").exists())
    qp.write_text('{"not": "a list"}', encoding="utf-8")
    check("ไฟล์คิวเพี้ยน (ไม่ใช่ลิสต์) → คิวว่าง ไม่พัง", engine.load_queue() == [])
    qp.write_text('[{"id":"x"}, 5, {"id":"y","path":"a.xlsx"}]', encoding="utf-8")
    check("งานเพี้ยนในคิวถูกกรองทิ้ง", [j["id"] for j in engine.load_queue()] == ["y"])

    print("\n── API ──")
    server.app.config["TESTING"] = True
    c = server.app.test_client()
    c.post("/api/setup", json={"username": "admin1", "password": "secret9",
                               "password2": "secret9", "display_name": "Admin"})
    done = xlsx("done.xlsx", [3101501909800, 3101501909801], [" - ", " - "])
    part = xlsx("part.xlsx", [3101501909800, 3101501909801], [" - ", "ERROR: x"])
    engine.save_queue([{"id": "d", "path": str(done), "user_id": 1},
                       {"id": "p", "path": str(part), "user_id": 1},
                       {"id": "m", "path": str(TEST_DATA / "gone.xlsx"), "user_id": 1}])
    q = {j["id"]: j for j in c.get("/api/queue").get_json()}
    check("/api/queue: ไฟล์ครบ → done", q["d"]["done"] is True and q["d"]["searched"] == 2)
    check("/api/queue: มีแถว ERROR → ไม่ done + บอกจำนวน",
          q["p"]["done"] is False and q["p"]["errors"] == 1 and q["p"]["pending"] == 0)
    check("/api/queue: ไฟล์หาย → missing", q["m"]["missing"] is True)

    calls = []
    real_busy, real_start = server.manager.is_busy, server.manager.start
    server.manager.start = lambda jobs, account, **kw: (calls.append((jobs, kw)), (True, "ok"))[1]
    r = c.post("/api/run/start", json={"account": ""})
    check("เริ่มรอบ: ส่งไปเฉพาะไฟล์ที่ยังไม่ครบ (ไม่รวมไฟล์ครบ/ไฟล์หาย)",
          r.status_code == 200 and [j["id"] for j in calls[-1][0]] == ["p"],
          str([j["id"] for j in calls[-1][0]]) if calls else r.get_data(as_text=True))
    r = c.post("/api/run/retry-errors", json={})
    check("ลองแถวผิดพลาด: เปิดเฉพาะไฟล์ที่มี ERROR", r.status_code == 200
          and [j["id"] for j in calls[-1][0]] == ["p"] and calls[-1][1].get("retry_only") is True)

    server.manager.is_busy = lambda: True
    r = c.post("/api/run/start", json={"account": ""})
    check("มีรอบค้าง → 409 พร้อมคีย์ error (หน้าเว็บแจ้งผู้ใช้ได้)",
          r.status_code == 409 and r.get_json().get("error"))
    r = c.post("/api/queue/prune")
    check("กำลังมีรอบทำงาน → ไม่แตะคิว", r.get_json().get("removed") == 0
          and len(engine.load_queue()) == 3)
    server.manager.is_busy = real_busy
    r = c.post("/api/queue/prune").get_json()
    check("prune: เอาไฟล์ครบ + ไฟล์หายออก เหลือไฟล์ที่ยังมีงาน",
          r["removed"] == 2 and [j["id"] for j in engine.load_queue()] == ["p"], str(r))
    engine.save_queue([{"id": "d", "path": str(done), "user_id": 1}])
    r = c.post("/api/run/retry-errors", json={})
    check("ไม่มีแถวผิดพลาดค้าง → แจ้งชัดเจน", r.status_code == 400 and "ไม่มีแถวที่ผิดพลาด" in r.get_json()["error"])
    server.manager.start = real_start

    print("\n── เพดานลองใหม่ต่อรอบ ──")
    import openpyxl
    f = xlsx("retry.xlsx", [3101501909800 + i for i in range(4)],
             ["ERROR#3: old", "ERROR!: หน้าค้นหา CRIMES เปลี่ยนไป", " - ", "ERROR: x"])
    wb = openpyxl.load_workbook(f)
    ws = wb.active
    m = worker.RunManager()
    ctx = {"id_col": 2, "out_col": 6, "key": str(f), "ids": {}}
    rows = m._collect_error_rows(ws, ctx, 2, None)
    check("รอบใหม่: ERROR#3 และ ERROR! จากรอบก่อนถูกเก็บมาลองใหม่ (เดิมข้ามตลอดไป)",
          rows == [2, 3, 5], str(rows))
    m._tries[(str(f), 2)] = m.RETRY_MAX_ATTEMPTS
    rows = m._collect_error_rows(ws, ctx, 2, None)
    check("ผิดครบเพดานในรอบนี้แล้ว → หยุดลองในรอบนี้ (กันเผาโควตา)", rows == [3, 5], str(rows))
    check("กดลองเองได้เสมอ (force)", m._collect_error_rows(ws, ctx, 2, None, force=True) == [2, 3, 5])
    check("นับแถวที่ยังค้างทั้งหมดได้ (บอกผู้ใช้)",
          len(m._collect_error_rows(ws, ctx, 2, None, all_rows=True)) == 3)
    t0 = time.time()
    m._stop = True
    check("หน่วงก่อนรอบลองใหม่กดหยุดได้ทันที", m._sleep_interruptible(30) is False and time.time() - t0 < 1)

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
