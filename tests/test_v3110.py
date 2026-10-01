#!/usr/bin/env python3
"""v3.11.0 — ยืนยันผลการค้นตามจำนวนแถวในไฟล์ + ส่งไฟล์ผลขึ้น Google Drive ผ่านศูนย์กลาง (Superadmin ตั้งค่าเท่านั้น)

ผู้ใช้สั่ง: "เพิ่มโมดุลเมื่อค้นหาเสร็จสมบูรณ์แล้ว ให้มีช่องทางในการบันทึกและยืนยันผลการค้น ตามจำนวนแถวในไฟล์ที่ค้น
            และช่องทางสำหรับอัพโหลดส่งไฟล์ที่สมบูรณ์ไว้ในกูเกิลไดร์ฟ ที่ทำการเชื่อมต่อในการตั้งค่า (Superadmin เป็นผู้ตั้งค่าเท่านั้น)"
  เครื่อง A = Super Admin (โปรเซสนี้) · เครื่อง B = เครื่องสมาชิก (โปรเซสแยก) · ศูนย์กลาง = hub_gas.js ตัวจริงบน Node + Drive จำลอง
  • สแกนไฟล์จริงจำแนก พบ/ไม่พบ/เลขใช้ไม่ได้/ผิดพลาด/ยังไม่ค้น → สรุปเข้า file_results เมื่อค้นจบรอบ
  • ยืนยันได้เฉพาะไฟล์ที่ครบทุกแถวและไม่มีแถวผิดพลาด · ตัวเลขเปลี่ยนหลังยืนยัน = ต้องยืนยันใหม่
  • ตั้งค่าไดรฟ์: ค่ากลางในชีต config — ตั้งได้เฉพาะก้อนที่เซ็น ADMIN_TOKEN (เครื่องผู้จัดการ) · ค่าตั้งต้นปิด
  • ส่งไฟล์: ต้องยืนยันก่อน · ศูนย์กลางปัดตกเมื่อปิดรับ/เกินโควตาวัน · ไฟล์ถึงไดรฟ์ครบไบต์ (sha256 ตรง) · ชื่อไทยคงเดิม
  • เครื่องสมาชิก (ไม่มีรหัสผู้ดูแล) ส่งไฟล์ได้ แต่ตั้งค่าไดรฟ์ไม่ได้
"""
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v3110_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "A")
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "A_up")
from _app import APP  # noqa: E402
import config  # noqa: E402

for _n in ("seed_account.json", config.HUB_SEED_FILE_NAME):
    (config.BASE_DIR / _n).unlink(missing_ok=True)

from backend import server  # noqa: E402
import db  # noqa: E402
import engine  # noqa: E402
import hub  # noqa: E402

HERE = Path(__file__).resolve().parent
STATE = TEST_DATA / "hub_sheet.json"
TOKEN = "test-hub-token-0123456789"
ADMIN = "test-admin-token-9876543210"
URL = "https://script.google.com/macros/s/TEST/exec"
FOUND1 = "คดีที่1 ลักทรัพย์ ปี2560"
PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def fake_post(url, payload, tok, admin_token=None, timeout=None):
    body = hub.encode_payload(payload).decode("ascii")
    param = {"sign": hub.sign(body.encode("utf-8"), tok)}
    if admin_token:
        param["asign"] = hub.sign(body.encode("utf-8"), admin_token)
    req = {"method": "POST", "body": body, "parameter": param}
    r = subprocess.run(["node", str(HERE / "gas_harness.js"), str(STATE), TOKEN, ADMIN],
                       input=json.dumps(req), capture_output=True, text=True, timeout=120)
    if r.returncode:
        raise RuntimeError("harness: " + r.stderr[-300:])
    return json.loads(r.stdout)


def machine_b(cmds, data="B"):
    env = dict(os.environ, CRIMES_DATA_DIR=str(TEST_DATA / data), CRIMES_UPLOAD_DIR=str(TEST_DATA / (data + "_up")),
               PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, str(HERE / "_hub_machine.py"), str(TEST_DATA / data), str(STATE), TOKEN, ADMIN],
                       input=json.dumps(cmds), env=env, capture_output=True, text=True, timeout=180)
    if r.returncode:
        raise RuntimeError("machine B: " + r.stderr[-800:])
    line = [ln for ln in r.stdout.splitlines() if ln.startswith("{")][-1]
    return json.loads(line)


def mk_id(seed, i):
    base = [int(d) for d in f"3{seed % 10}{i:010d}"]
    chk = (11 - sum(d * (13 - j) for j, d in enumerate(base)) % 11) % 10
    return "".join(map(str, base)) + str(chk)


def make_xlsx(name, n, results=(), seed=0):
    """ไฟล์งานจริง: หัวตาราง + เลขบัตร 13 หลักคอลัมน์ B · ผล (ถ้ามี) คอลัมน์ F"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="ลำดับ")
    ws.cell(row=1, column=2, value="เลขบัตรประชาชน")
    for i in range(n):
        ws.cell(row=2 + i, column=1, value=i + 1)
        ws.cell(row=2 + i, column=2, value=mk_id(seed, i))
        if i < len(results) and results[i]:
            ws.cell(row=2 + i, column=6, value=results[i])
    src = TEST_DATA / "srcA"
    src.mkdir(exist_ok=True)
    fp = src / name
    wb.save(fp)
    return fp


def upload(c, fp):
    with open(fp, "rb") as fh:
        r = c.post("/api/upload", data={"file": (fh, Path(fp).name)}, content_type="multipart/form-data")
    return r.status_code, r.get_json()


def fill(path, values):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for i, v in enumerate(values):
        if v:
            ws.cell(row=2 + i, column=6, value=v)
    wb.save(path)


def fr_rows(c):
    return (c.get("/api/files/results").get_json() or {}).get("rows", [])


def fr_by_file(c, name):
    return next((x for x in fr_rows(c) if x["file"] == name), None)


def sheet(name):
    st = json.loads(STATE.read_text(encoding="utf-8"))
    return st.get("sheets", {}).get(name, {}).get("rows", [])


def drive_files():
    st = json.loads(STATE.read_text(encoding="utf-8"))
    out = []
    for fid, f in (st.get("drive", {}).get("folders", {}) or {}).items():
        for file in f["files"]:
            out.append(dict(file, folder=f["name"], folder_id=fid))
    return out


def main():
    check("มี Node สำหรับรันสคริปต์ศูนย์กลางตัวจริง", shutil.which("node") is not None)
    server.hub._post = fake_post
    server._hub_kick = lambda *a, **k: None
    server.app.config["TESTING"] = True
    c = server.app.test_client()
    code = hub.make_connect_code(URL, TOKEN)

    print("── จำแนกผลจากไฟล์จริง (engine.scan_progress) ──")
    fp = make_xlsx("จำแนก.xlsx", 6, [FOUND1, " - ", "⚠ เลขบัตรไม่ถูกต้อง (สั้นไป) — ไม่ได้ค้น",
                                     "ERROR: เว็บช้า", "", "คดีที่1 บุกรุก ปี2561/ คดีที่2 ลักทรัพย์ ปี2559"])
    pr = engine.scan_progress(str(fp), "B", "F")
    check("นับครบ: 6 แถว · มีผล 4 · พบ 2 · ไม่พบ 1 · เลขใช้ไม่ได้ 1 · ผิดพลาด 1 · ยังไม่ค้น 1",
          (pr["total"], pr["ok"], pr["found"], pr["notfound"], pr["invalid_marked"], pr["errors"], pr["pending"])
          == (6, 4, 2, 1, 1, 1, 1), str(pr))
    check("ตัวนับใหม่ไม่ทับคีย์ 'invalid' ของหน้าคิว (ป้าย 'ต้องแก้ไฟล์' ต้องไม่โผล่กับไฟล์ปกติ)",
          "invalid" not in pr)

    print("\n── เครื่อง A: Super Admin + สมาชิก + เชื่อมศูนย์กลาง ──")
    r = c.post("/api/setup", json={"username": "admin1", "password": "secret9", "password2": "secret9",
                                   "display_name": "แอดมิน", "hub_code": code, "hub_admin_token": ADMIN})
    check("ตั้งเครื่องแรก (Super Admin + รหัสผู้ดูแล)", r.status_code == 200)
    c.post("/api/admin/members", json={"username": "worker1", "password": "pass66", "display_name": "คนงานหนึ่ง",
                                       "role": "member"})
    check("(เตรียม) สมาชิกขึ้นกลาง", server._members_sync_now("test").get("ok") and db.count_pending_members() == 0)

    print("\n── สรุปผลรายไฟล์: เข้าเอง · นับตามไฟล์จริง · ยืนยันเมื่อครบเท่านั้น ──")
    st1, up1 = upload(c, make_xlsx("งานเดือนตุลา.xlsx", 4, seed=1))
    check("อัปโหลดไฟล์ 4 แถวเข้าคิว", st1 == 200 and up1.get("rows") == 4, str(up1))
    path1 = up1["job"]["path"]
    check("ไฟล์ยังไม่ครบ → ไม่มีสรุปในรายการ", fr_by_file(c, "งานเดือนตุลา.xlsx") is None)
    fill(path1, [FOUND1, " - ", "ERROR: หน้าเว็บไม่ตอบ", None])
    check("ค้นแล้ว 3/4 (มี ERROR) → ยังไม่เข้ารายการ (ยังมีแถวไม่ได้ค้น)", fr_by_file(c, "งานเดือนตุลา.xlsx") is None)
    fill(path1, [None, None, None, " - "])
    row = fr_by_file(c, "งานเดือนตุลา.xlsx")
    check("ครบทุกแถว (เหลือ ERROR 1) → เข้ารายการ · done=False",
          row and row["rows_total"] == 4 and row["errors"] == 1 and row["pending"] == 0 and not row["done"], str(row))
    r = c.post(f"/api/files/results/{row['id']}/confirm", json={})
    check("ยืนยันไฟล์ที่ยังมีแถวผิดพลาด → 400 พร้อมตัวเลข", r.status_code == 400 and "ผิดพลาด" in r.get_json()["error"])
    fill(path1, [None, None, "คดีที่1 ฉ้อโกง ปี2562", None])
    row = fr_by_file(c, "งานเดือนตุลา.xlsx")
    check("แก้แถว ERROR เป็นผลจริง → done=True · พบ 2 · ไม่พบ 2",
          row and row["done"] and row["found"] == 2 and row["notfound"] == 2 and row["errors"] == 0, str(row))
    r = c.post(f"/api/files/results/{row['id']}/confirm", json={"note": "ตรวจแล้ว"})
    j = r.get_json()
    check("ยืนยันสำเร็จ · จดคนยืนยัน", r.status_code == 200 and j["row"]["confirmed"]
          and j["row"]["confirmed_name"] == "แอดมิน", str(j))
    rid1 = row["id"]

    print("\n── ตัวเลขเปลี่ยนหลังยืนยัน = การยืนยันถูกล้าง ──")
    fill(path1, [None, "ERROR: ลองใหม่", None, None])
    row = fr_by_file(c, "งานเดือนตุลา.xlsx")
    check("ไฟล์ถูกค้นใหม่/ผลเปลี่ยน → สรุปใหม่ และต้องยืนยันใหม่",
          row and row["id"] == rid1 and not row["confirmed"] and row["errors"] == 1 and not row["done"], str(row))
    fill(path1, [None, " - ", None, None])
    r = c.post(f"/api/files/results/{rid1}/confirm", json={})
    check("กลับมาครบ → ยืนยันรอบใหม่ได้", r.status_code == 200 and r.get_json()["row"]["confirmed"])

    print("\n── คิวถูก prune → สรุปถูกเก็บก่อนงานหาย ──")
    st2, up2 = upload(c, make_xlsx("งานชุดสอง.xlsx", 2, seed=2))
    path2 = up2["job"]["path"]
    fill(path2, [FOUND1, "⚠ เลขบัตรไม่ถูกต้อง (สั้นไป) — ไม่ได้ค้น"])
    qj = next(j for j in c.get("/api/queue").get_json() if j["path"] == path2)
    check("ไฟล์ครบที่มีแถว 'เลขใช้ไม่ได้' → คิวยังขึ้น File ready (ไม่ใช่ป้าย 'ต้องแก้ไฟล์' — บั๊กชนคีย์ invalid)",
          qj["done"] and qj["invalid"] == "", str(qj))
    fill(path2, [None, " - "])
    r = c.post("/api/queue/prune")
    names = r.get_json().get("names", [])
    check("prune เอาไฟล์ที่ครบออกจากคิว", "งานชุดสอง.xlsx" in names, str(names))
    row2 = fr_by_file(c, "งานชุดสอง.xlsx")
    check("สรุปยังอยู่หลังงานหายจากคิว · exists=True", row2 and row2["done"] and row2["exists"], str(row2))

    print("\n── ตั้งค่า Google Drive: Superadmin + เครื่องผู้จัดการเท่านั้น · ค่าตั้งต้นปิด ──")
    r = c.get("/api/hub/drive?refresh=1")
    d = r.get_json()["drive"]
    check("สถานะเริ่มต้น: ตรวจได้ · ปิดอยู่ · โฟลเดอร์ค่าตั้งต้น",
          d.get("ok") and not d.get("enabled") and d.get("folder") == "CRIMES AUTO ไฟล์ผลการค้น", str(d))
    r = c.post(f"/api/files/results/{rid1}/drive", json={})
    check("ส่งไฟล์ตอนศูนย์กลางยังปิดรับ → ศูนย์กลางปัดตก (บอกวิธีเปิด)",
          r.status_code == 502 and "ยังไม่เปิดรับ" in r.get_json()["error"], str(r.get_json()))
    tok_backup = server.auth.load_config().get("hub_admin_token")
    server.auth.update_config(hub_admin_token="")
    r = c.post("/api/hub/drive", json={"enabled": True, "folder": "ผลงานหน่วย ก"})
    check("เครื่องไม่มีรหัสผู้ดูแล → ตั้งค่าไม่ได้ (need_admin_token)",
          r.status_code == 403 and r.get_json().get("need_admin_token"), str(r.get_json()))
    server.auth.update_config(hub_admin_token=tok_backup)
    r = c.post("/api/hub/drive", json={"enabled": True, "folder": "ผลงานหน่วย ก"})
    j = r.get_json()
    check("เครื่องผู้จัดการเปิดรับ + ตั้งชื่อโฟลเดอร์ → สำเร็จ พร้อมลิงก์โฟลเดอร์",
          r.status_code == 200 and j["drive"]["enabled"] and j["drive"]["folder"] == "ผลงานหน่วย ก"
          and j["drive"].get("folder_url", "").startswith("https://drive.google.com/drive/folders/"), str(j))

    print("\n── ส่งไฟล์ขึ้นไดรฟ์ (เครื่อง A) ──")
    r = c.post(f"/api/files/results/{row2['id']}/drive", json={})
    check("ไฟล์ที่ยังไม่ยืนยัน → 400 ต้องยืนยันก่อน", r.status_code == 400 and "ยืนยัน" in r.get_json()["error"])
    r = c.post(f"/api/files/results/{rid1}/drive", json={})
    j = r.get_json()
    check("ไฟล์ที่ยืนยันแล้ว → ส่งสำเร็จ ได้ลิงก์ไฟล์",
          r.status_code == 200 and j.get("url", "").startswith("https://drive.google.com/file/d/")
          and j["row"]["drive_url"] == j["url"], str(j))
    local_sha = hashlib.sha256(Path(path1).read_bytes()).hexdigest()
    df = drive_files()
    check("ไฟล์ถึงไดรฟ์ครบไบต์ (sha256 ตรงกับไฟล์ในเครื่อง) · ลงโฟลเดอร์ที่ตั้งไว้ · ชื่อมีจำนวนชื่อ",
          len(df) == 1 and df[0]["sha256"] == local_sha and df[0]["folder"] == "ผลงานหน่วย ก"
          and df[0]["name"] == "งานเดือนตุลา(4ชื่อ).xlsx", str(df))
    up_rows = sheet("uploads")
    check("ชีต uploads จดรายการ: ผู้ส่ง · เดือน · จำนวนแถว · ขนาด · sha256",
          len(up_rows) == 2 and up_rows[1][2] == "แอดมิน" and up_rows[1][5] == 4
          and up_rows[1][9] == Path(path1).stat().st_size and up_rows[1][10] == local_sha, str(up_rows[1:]))

    print("\n── ขอบเขตสิทธิ์ในเครื่อง: สมาชิกเห็น/จัดการเฉพาะของตัวเอง ──")
    cw = server.app.test_client()
    r = cw.post("/api/login", json={"username": "worker1", "password": "pass66"})
    check("worker1 เข้าระบบ", r.status_code == 200)
    check("worker1 ไม่เห็นสรุปของแอดมิน", fr_rows(cw) == [])
    r = cw.post(f"/api/files/results/{rid1}/confirm", json={})
    check("worker1 ยืนยันงานของคนอื่น → 403", r.status_code == 403)
    r = cw.post("/api/hub/drive", json={"enabled": False})
    check("worker1 ตั้งค่าไดรฟ์ → 403 (ต้องเป็น Super Admin)", r.status_code == 403)
    r = cw.get("/api/hub/drive")
    dw = r.get_json()["drive"]
    check("worker1 เห็นสถานะเปิดรับ แต่ไม่เห็นลิงก์โฟลเดอร์", dw.get("enabled") and "folder_url" not in dw, str(dw))

    print("\n── เครื่อง B (สมาชิก ไม่มีรหัสผู้ดูแล): ทำไฟล์ชื่อไทย → ยืนยัน → ส่งขึ้นไดรฟ์ ──")
    out = machine_b([
        {"op": "join", "code": code},
        {"op": "login", "u": "worker1", "p": "pass66"},
        {"op": "make_file", "n": 3, "name": "งานทดสอบ ๑.xlsx", "seed": 7, "as": "mk"},
    ])
    check("B เข้าร่วมองค์กร + worker1 เข้าระบบ", out["join"][0] == 200 and out["login"][0] == 200, str(out))
    out = machine_b([
        {"op": "login", "u": "worker1", "p": "pass66"},
        {"op": "upload_api", "path": out["mk"], "as": "up"},
        {"op": "get", "path": "/api/queue", "as": "q"},
    ])
    check("B อัปโหลดไฟล์ 3 แถวเข้าคิว", out["up"][0] == 200 and out["up"][1]["rows"] == 3, str(out["up"]))
    bpath = out["up"][1]["job"]["path"]
    out = machine_b([
        {"op": "login", "u": "worker1", "p": "pass66"},
        {"op": "fill_results", "path": bpath, "values": [FOUND1, " - ", " - "]},
        {"op": "get", "path": "/api/files/results", "as": "fr"},
    ])
    brow = next(x for x in out["fr"][1]["rows"] if x["file"] == "งานทดสอบ ๑.xlsx")
    check("B: สรุปครบ 3 แถว (พบ 1 ไม่พบ 2) · เห็นสถานะไดรฟ์เปิดจากค่ากลาง",
          brow["done"] and brow["found"] == 1 and brow["notfound"] == 2 and brow["can_act"], str(brow))
    out = machine_b([
        {"op": "login", "u": "worker1", "p": "pass66"},
        {"op": "post", "path": f"/api/files/results/{brow['id']}/confirm", "json": {}, "as": "cf"},
        {"op": "post", "path": f"/api/files/results/{brow['id']}/drive", "json": {}, "as": "dv"},
    ])
    check("B ยืนยัน + ส่งขึ้นไดรฟ์สำเร็จ (ใช้แค่รหัสร่วม)",
          out["cf"][0] == 200 and out["dv"][0] == 200
          and out["dv"][1]["url"].startswith("https://drive.google.com/file/d/"), str(out))
    df = drive_files()
    bfile = next(x for x in df if x["name"].startswith("งานทดสอบ ๑"))
    check("ไฟล์ชื่อไทยของ B ถึงไดรฟ์ · ชื่อคงเดิม (เติมจำนวนชื่อ)", bfile["name"] == "งานทดสอบ ๑(3ชื่อ).xlsx", str(df))
    out = machine_b([{"op": "set_admin_token", "token": ""},
                     {"op": "login", "u": "worker1", "p": "pass66"},
                     {"op": "drive_direct", "set": {"enabled": False}, "as": "ds"}])
    check("B (ไม่มีรหัสผู้ดูแล) ตั้งค่าไดรฟ์ตรงถึงศูนย์กลาง → ศูนย์กลางปัดตก",
          out["ds"].get("ok") is False and "ผู้ดูแล" in out["ds"].get("error", ""), str(out["ds"]))

    print("\n── เพดานศูนย์กลาง: โควตารายวัน · ปิดรับแล้วปัดตก ──")
    st = json.loads(STATE.read_text(encoding="utf-8"))
    today_iso = up_rows[1][0]
    st["sheets"]["uploads"]["rows"] += [[today_iso, "x", "x", f"junk{i}.xlsx", "2026-10", 1, 0, 1, 0, 9, "s", "f", "u"]
                                        for i in range(40)]
    STATE.write_text(json.dumps(st), encoding="utf-8")
    res = hub.upload_file(URL, TOKEN, {"name": "เกินโควตา.xlsx", "ym": "2026-10", "user": "x",
                                       "rows_total": 1, "found": 0, "notfound": 1, "errors": 0}, b"PK\x03\x04x")
    check("เกินโควตารายวัน → ศูนย์กลางปัดตก", res.get("ok") is False and "ครบ" in res.get("error", ""), str(res))
    st["sheets"]["uploads"]["rows"] = st["sheets"]["uploads"]["rows"][:3]
    STATE.write_text(json.dumps(st), encoding="utf-8")

    print("\n── สคริปต์รุ่นเก่า: โปรแกรมบอกให้วางสคริปต์ 3.11.0 (ไม่เงียบ) ──")
    real_post = hub._post
    hub._post = lambda *a, **k: {"ok": True, "ver": "3.10.2", "stored": 0, "kind": "upload"}
    res = hub.upload_file(URL, TOKEN, {"name": "x.xlsx", "ym": "2026-10", "user": "x",
                                       "rows_total": 1, "found": 1, "notfound": 0, "errors": 0}, b"PK")
    check("upload กับสคริปต์เก่า → แจ้งวิธีวางสคริปต์ 3.11.0",
          res.get("ok") is False and "3.11.0" in res.get("error", ""), str(res))
    res = hub.drive_config(URL, TOKEN)
    check("drive_cfg กับสคริปต์เก่า → แจ้งวิธีวางสคริปต์ 3.11.0",
          res.get("ok") is False and "3.11.0" in res.get("error", ""), str(res))
    hub._post = real_post

    print("\n── ความปลอดภัยฝั่งโปรแกรม: เพดานขนาด · ลบรายการสรุป ──")
    res = hub.upload_file(URL, TOKEN, {"name": "ใหญ่.xlsx", "ym": "2026-10", "user": "x",
                                       "rows_total": 1, "found": 0, "notfound": 1, "errors": 0},
                          b"x" * (hub.MAX_UPLOAD_MB * 1024 * 1024 + 1))
    check(f"ไฟล์เกิน {hub.MAX_UPLOAD_MB} MB → โปรแกรมไม่ส่ง", res.get("ok") is False and "ใหญ่เกิน" in res["error"])
    r = c.delete(f"/api/files/results/{row2['id']}")
    check("ลบรายการสรุป (ไม่แตะไฟล์)", r.status_code == 200 and fr_by_file(c, "งานชุดสอง.xlsx") is None
          and Path(path2).exists())

    print("\n── แหล่งที่มา/กติกาที่ต้องคงไว้ ──")
    gas_src = (HERE.parent / "tools" / "hub_gas.js").read_text(encoding="utf-8")
    app_backend = Path(hub.__file__).resolve().parent
    hub_src = (app_backend / "hub.py").read_text(encoding="utf-8")
    worker_src = (app_backend / "worker.py").read_text(encoding="utf-8")
    html_src = (app_backend.parent / "frontend" / "index.html").read_text(encoding="utf-8")
    check("hub_gas.js = 3.11.0 · มี upload/drive_cfg/โควตา", "HUB_SCRIPT_VERSION = '3.11.0'" in gas_src
          and "kind === 'upload'" in gas_src and "kind === 'drive_cfg'" in gas_src
          and "MAX_UPLOADS_PER_DAY" in gas_src)
    check("โปรแกรมยังต้องการสคริปต์ ≥ 3.10.0 เท่าเดิม (ไดรฟ์เป็นของเสริม ไม่บังคับวางสคริปต์ใหม่)",
          hub.HUB_SCRIPT_REQUIRED == "3.10.0")
    check("กติกาความเป็นส่วนตัวใน hub.py จดข้อยกเว้น v3.11.0 ชัดเจน (กดส่งเองเท่านั้น)",
          "ข้อยกเว้นเดียว (v3.11.0" in hub_src and "กดปุ่มส่งเองทีละไฟล์" in hub_src)
    check("worker เก็บสรุปผลเมื่อค้นจบไฟล์", "upsert_file_result" in worker_src)
    check("หน้าเว็บมีการ์ดยืนยันผล + ตั้งค่าไดรฟ์ (Superadmin)", "fileResultsCard" in html_src
          and "btnHubDriveSave" in html_src and "fr-confirm" in html_src and "loadDriveCfg" in html_src)

    print(f"\n==== ผล: ผ่าน {PASS} · ตก {FAIL} ====")
    return 1 if FAIL else 0


if __name__ == "__main__":
    code = main()
    shutil.rmtree(TEST_DATA, ignore_errors=True)
    sys.exit(code)
