#!/usr/bin/env python3
"""v3.1.0 — ทดสอบ backend: กองกลาง (ยอดรวมทุกคนเห็น) · Balance ตามขอบเขต+กราฟ ·
เหตุผล/ความเร็ว/ขั้นตอน (ของเดิมต้องไม่พัง)"""
import os
import shutil
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v310_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA)
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "uploads")

from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))

from backend import db, server  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def seed(uid, account, n, days_ago=0):
    ts = datetime.now() - timedelta(days=days_ago)
    with db.get_conn() as c:
        for i in range(n):
            c.execute(
                "INSERT INTO searches(ts,ym,account,file,row,national_id,outcome,"
                "case_count,detail,user_id,id_hash,incomplete,attempt) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (ts.isoformat(timespec="seconds"), ts.strftime("%Y-%m"), account,
                 "s.xlsx", i + 2, "1-23xx-xxxxx-xx-1", "notfound", 0, "", uid,
                 f"u{uid}d{days_ago}i{i}", 0, 1))


def main():
    server.app.config["TESTING"] = True
    admin = server.app.test_client()
    m1 = server.app.test_client()     # สมาชิกในทีม
    m2 = server.app.test_client()     # สมาชิกนอกทีม

    admin.post("/api/setup", json={"username": "admin1", "password": "secret9",
                                   "password2": "secret9", "display_name": "Admin"})
    admin.post("/api/admin/members", json={"username": "somchai", "password": "pass66",
                                           "display_name": "สมชาย", "role": "member"})
    admin.post("/api/admin/members", json={"username": "wichai", "password": "pass66",
                                           "display_name": "วิชัย", "role": "member"})
    m1.post("/api/login", json={"username": "somchai", "password": "pass66"})
    m2.post("/api/login", json={"username": "wichai", "password": "pass66"})
    users = {u["username"]: u["id"] for u in db.list_users()}

    seed(users["admin1"], "acctA", 5)
    seed(users["somchai"], "acctB", 3)
    seed(users["somchai"], "acctB", 2, days_ago=1)
    seed(users["wichai"], "acctC", 7)
    with db.get_conn() as c:   # รายการรับ/จ่ายกองกลาง
        now = datetime.now()
        c.execute("INSERT INTO ledger(ts,ym,user_id,kind,amount,note,created_by)"
                  " VALUES(?,?,?,?,?,?,?)",
                  (now.isoformat(timespec="seconds"), now.strftime("%Y-%m"),
                   users["somchai"], "in", 100.0, "", 1))
        c.execute("INSERT INTO ledger(ts,ym,user_id,kind,amount,note,created_by)"
                  " VALUES(?,?,?,?,?,?,?)",
                  (now.isoformat(timespec="seconds"), now.strftime("%Y-%m"),
                   users["wichai"], "out", 40.0, "", 1))

    # ── กองกลาง: ทุกคนเห็นยอดรวมชุดเดียวกัน ──
    a = admin.get("/api/live/board").get_json()
    b = m2.get("/api/live/board").get_json()
    check("สมาชิกธรรมดาเปิดกองกลางได้", "totals" in b, str(b)[:120])
    # v3.2.0: กองกลาง = สมุดบันทึกรับ/จ่ายเท่านั้น (รับ 100 จ่าย 40) — ไม่รวมผลค้น 17 รายการ
    check("ยอดรวมกองกลางเท่ากันทุกคน (ไม่ใช่ของคนใดคนหนึ่ง) และเป็นสมุดบันทึกล้วน",
          a["totals"] == b["totals"] and a["totals"] == {"in": 100.0, "out": 40.0, "net": 60.0, "entries": 2},
          f'{a["totals"]} vs {b["totals"]}')
    check("มีข้อมูลกราฟรายวัน (daily)", len(a.get("daily") or []) >= 1)
    check("สมาชิกไม่มีสิทธิ์ → ไม่เห็นรายคน/รายการ",
          b["rows"] == [] and b["entries"] == [] and b["detail"] is False)
    check("แอดมินเห็นรายคนที่มีรายการบันทึก (2 คน)", len(a["rows"]) == 2 and a["detail"] is True)

    # ── Balance ตามขอบเขต + Real time ──
    s_admin = admin.get("/api/admin/billing/summary").get_json()
    check("แอดมินเห็นทุกคน + จัดการได้",
          len(s_admin["members"]) == 3 and s_admin["can_manage"] is True)
    check("Balance มีข้อมูลกราฟ daily", len(s_admin.get("daily") or []) >= 2)

    s_m2 = m2.get("/api/admin/billing/summary").get_json()
    check("สมาชิกนอกทีมเห็นเฉพาะตัวเอง",
          len(s_m2["members"]) == 1 and s_m2["members"][0]["user_id"] == users["wichai"]
          and s_m2["can_manage"] is False, str(s_m2["members"]))
    check("ยอดรวมเป็นของตัวเองเท่านั้น", s_m2["total_count"] == 7)

    # ตั้งทีม: admin1+somchai อยู่ทีมเดียวกัน แล้วให้ somchai มีสิทธิ์ดูทีม
    r = admin.post("/api/admin/teams", json={"name": "ทีมสืบ",
                                             "member_ids": [users["admin1"], users["somchai"]]})
    check("สร้างทีม", r.status_code == 200, r.get_data(as_text=True))
    r = admin.put(f"/api/admin/members/{users['somchai']}",
                  json={"permissions": {"view_team": True}})
    check("ให้สิทธิ์ดูทีมแก่สมชาย", r.status_code == 200)
    s_m1 = m1.get("/api/admin/billing/summary").get_json()
    ids = sorted(x["user_id"] for x in s_m1["members"])
    check("สมาชิกมีสิทธิ์ดูทีม → เห็นทั้งทีม (ตัวเอง+แอดมินร่วมทีม)",
          ids == sorted([users["admin1"], users["somchai"]]), str(ids))

    # detail ตามขอบเขต
    r = m1.get(f"/api/admin/billing/detail/{users['admin1']}")
    check("ดูรายวันของเพื่อนร่วมทีมได้", r.status_code == 200)
    r = m1.get(f"/api/admin/billing/detail/{users['wichai']}")
    check("ดูรายวันของคนนอกทีมไม่ได้ (403)", r.status_code == 403)
    r = m2.get(f"/api/admin/billing/detail/{users['wichai']}")
    check("ดูรายวันของตัวเองได้", r.status_code == 200)

    # ── ของเดิมต้องไม่พัง: เหตุผลตอนกดเริ่ม + ความเร็ว ──
    from backend import engine
    import openpyxl
    wb = openpyxl.Workbook()           # v3.5.0: สถานะไฟล์อ่านจากไฟล์จริง → ต้องมีไฟล์อยู่จริง
    wb.active.cell(1, 2, "เลขบัตร")
    wb.active.cell(2, 2, 3101501909804)
    fpath = TEST_DATA / "a.xlsx"
    wb.save(fpath)
    engine.save_queue([{"id": "j1", "path": str(fpath), "user_id": users["admin1"]}])
    calls = []
    server.manager.is_busy = lambda: False
    server.manager.start = lambda jobs, account, **kw: (calls.append(jobs), (True, "ok"))[1]
    r = admin.post("/api/run/start", json={"account": "", "reason": "ตรวจประวัติ"})
    check("ส่ง reason ตอนกดเริ่มยังทำงาน", r.status_code == 200
          and engine.load_queue()[0].get("reason") == "ตรวจประวัติ")
    r = admin.post("/api/run/speed", json={"level": 4})
    check("ตั้งความเร็วด้วย level ได้", r.status_code == 200
          and r.get_json().get("speed") == 4, r.get_data(as_text=True))

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
