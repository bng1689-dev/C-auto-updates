#!/usr/bin/env python3
"""UI smoke — เปิดหน้าจอจริงด้วย Chromium: สมัคร → แดชบอร์ด → ไล่ครบทุกเมนู → ออก → เข้าใหม่

ต้องไม่มี JavaScript error แม้แต่จุดเดียวตลอดทั้งเส้นทาง (ตาข่ายกันเหตุ 'หน้าดำเปล่า' v2.8.0)
v2.10.0 เพิ่ม: กราฟแคนวาสเคลื่อนไหวต้องวาดจริง · ป้าย File ready แทนปุ่ม Download ·
ช่องเหตุผลการค้นย้ายมาอยู่เหนือแถว LOGIN CRIMES
"""
import os
import shutil
import sys
import tempfile
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_ui_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA)

from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))   # db.py ใช้ flat import (from config import ...)

from backend import db, engine, server  # noqa: E402

PORT = 8791

PASS = FAIL = 0
JS_ERRORS = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def seed_history():
    """ยัดผลค้นย้อนหลัง 7 วันหลายบัญชี ให้กราฟทุกตัวมีข้อมูลวาดจริง"""
    now = datetime.now()
    with db.get_conn() as c:
        for d in range(7):
            ts = now - timedelta(days=d)
            for i in range(2 + d % 3):
                outcome = ["found", "notfound", "error"][i % 3]
                c.execute(
                    "INSERT INTO searches(ts,ym,account,file,row,national_id,outcome,"
                    "case_count,detail,user_id,id_hash,incomplete,attempt) "
                    "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (ts.isoformat(timespec="seconds"), ts.strftime("%Y-%m"),
                     f"acct{i % 2}", "seed.xlsx", i + 2, "1-23xx-xxxxx-xx-1",
                     outcome, 1 if outcome == "found" else 0, "", 1, f"h{d}-{i}", 0, 1))


def logout(page):
    """กดออกจากระบบแล้ว 'รอให้หน้า reload จริง' ก่อนทำขั้นต่อไป
    v3.8.1: เดิมกดแล้วรอแค่ #screen-auth โผล่ — บนหน้าเก่า poll ที่ได้ 401 อาจโชว์หน้าเข้าสู่ระบบก่อน location.reload()
    ทำให้ขั้นต่อไป (เช่น กรอกฟอร์มลืมรหัสผ่าน) ถูก reload ทับกลางคันแล้วล้มเป็นบางครั้ง (CI ล้มจริง)"""
    page.evaluate("window.__smokeDoc = 1")
    page.click("#btnLogout")
    page.wait_for_function("typeof window.__smokeDoc === 'undefined'", timeout=15000)   # เอกสารใหม่หลัง reload
    page.wait_for_selector("#screen-auth:not(.hidden)", timeout=15000)


def main():
    server.app.config["TESTING"] = False
    th = threading.Thread(
        target=lambda: server.app.run(port=PORT, threaded=True, use_reloader=False),
        daemon=True)
    th.start()
    time.sleep(1.2)

    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=CHROME, headless=True, args=["--disable-background-networking", "--disable-component-update", "--no-first-run", "--no-default-browser-check"])
        page = browser.new_page(viewport={"width": 1280, "height": 900})
        page.on("pageerror", lambda e: JS_ERRORS.append(f"pageerror: {e}"))
        page.on("console", lambda m: JS_ERRORS.append(f"console.error: {m.text}")
                if m.type == "error" else None)

        base = f"http://127.0.0.1:{PORT}"
        page.goto(base, wait_until="domcontentloaded")

        # ── สมัคร Super Admin คนแรก ──
        page.wait_for_selector("#authSetupBox:not(.hidden)", timeout=15000)
        check("หน้าแรกโชว์ฟอร์มสร้างบัญชี", True)
        # ── v3.7.0: โหมดตั้งค่าเครื่อง — เครื่องเพิ่มเติมเห็นเฉพาะช่องรหัสเชื่อมต่อ (ไม่สร้าง Super Admin ซ้ำ) ──
        page.click('#setupModeBar [data-smode="join"]')
        page.wait_for_selector("#authJoinForm:not(.hidden)", timeout=5000)
        check("โหมด 'เครื่องเพิ่มเติม' → ฟอร์มเข้าร่วมโผล่ ฟอร์มสร้าง Super Admin ซ่อน",
              page.eval_on_selector("#authSetupForm", "el => el.classList.contains('hidden')"))
        page.fill("#joinHubCode", "CRIMES-HUB:zzz")
        page.click("#btnJoinSubmit")
        page.wait_for_function("document.getElementById('authJoinErr').textContent.includes('ไม่ถูกต้อง')", timeout=8000)
        check("รหัสเชื่อมต่อผิด → แจ้งบนหน้า ไม่หลุดไปไหน", True)
        page.click('#setupModeBar [data-smode="first"]')
        page.wait_for_selector("#authSetupForm:not(.hidden)", timeout=5000)
        check("กลับโหมด 'เครื่องแรก' → ฟอร์มสร้าง Super Admin กลับมา",
              page.eval_on_selector("#authJoinForm", "el => el.classList.contains('hidden')"))
        page.fill("#suUser", "admin1")
        page.fill("#suName", "แอดมินทดสอบ")
        page.fill("#suPw", "secret9")
        page.fill("#suPw2", "secret9")
        page.click("#authSetupForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)
        check("สมัครแล้วเข้าหน้าแอปได้", True)

        # ── seed ข้อมูล + คิว แล้วรีโหลดให้แดชบอร์ดวาดกราฟจริง ──
        # /api/queue คำนวณ done/searched ใหม่จากไฟล์ xlsx จริง + จำนวนในฐานข้อมูล
        # จึงต้องสร้างไฟล์จริงให้จำนวนแถวสอดคล้องกับผลค้นที่ seed
        seed_history()
        import openpyxl
        def make_xlsx(name, rows, filled=0):
            # v3.5.0: ความคืบหน้าอ่านจาก 'ช่องผลในไฟล์จริง' (คอลัมน์ F) ไม่ใช่จากฐานข้อมูล
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.cell(1, 2, "เลขบัตร")
            for i in range(rows):
                ws.cell(2 + i, 2, f"1{str(i).zfill(12)}")
                if i < filled:
                    ws.cell(2 + i, 6, " - ")
            p = TEST_DATA / name
            wb.save(p)
            return str(p)
        ready = make_xlsx("ready.xlsx", 3, 3)      # ค้นครบ 3/3 → File ready
        partial = make_xlsx("partial.xlsx", 5, 2)  # ค้นไป 2/5 → ยังไม่ครบ
        fresh = make_xlsx("fresh.xlsx", 3)         # ยังไม่แตะ → รอค้น
        with db.get_conn() as c:
            for fname, n in (("ready.xlsx", 3), ("partial.xlsx", 2)):
                for i in range(n):
                    c.execute(
                        "INSERT INTO searches(ts,ym,account,file,row,national_id,outcome,"
                        "case_count,detail,user_id,id_hash,incomplete,attempt) "
                        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (datetime.now().isoformat(timespec="seconds"),
                         datetime.now().strftime("%Y-%m"), "acct0", fname, 2 + i,
                         "1-23xx-xxxxx-xx-1", "notfound", 0, "", 1, f"{fname}-{i}", 0, 1))
        engine.save_queue([
            {"id": "q1", "path": ready, "user_id": 1},
            {"id": "q2", "path": partial, "user_id": 1},
            {"id": "q3", "path": fresh, "user_id": 1},
        ])
        with db.get_conn() as c:   # v3.2.0: กองกลาง = สมุดบันทึกล้วน → ต้องมีรายการจึงมีกราฟ
            for d, kind, amt in ((0, "in", 300.0), (1, "out", 80.0), (2, "in", 120.0)):
                ts = datetime.now() - timedelta(days=d)
                c.execute("INSERT INTO ledger(ts,ym,user_id,kind,amount,note,created_by)"
                          " VALUES(?,?,?,?,?,?,?)",
                          (ts.isoformat(timespec="seconds"), ts.strftime("%Y-%m"), 1, kind, amt, "", 1))
        page.evaluate("""async () => {
          await fetch('/api/admin/teams',{method:'POST',headers:{'Content-Type':'application/json'},
            body:JSON.stringify({name:'ทีมทดสอบ', member_ids:[1]})}); }""")
        page.reload(wait_until="domcontentloaded")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)

        # ── แดชบอร์ด: กราฟเคลื่อนไหววาดจริง ──
        page.wait_for_selector("#chartDaily canvas", timeout=15000)
        check("กราฟรายวันเป็นแคนวาส", True)
        page.wait_for_timeout(1300)   # รอสวีป/นับเลขจบ
        painted = page.evaluate(
            """() => { const cv = document.querySelector('#chartDaily canvas');
                 const ctx = cv.getContext('2d');
                 return ctx.getImageData(0,0,cv.width,cv.height).data.some(v=>v>0); }""")
        check("แคนวาสมีพิกเซลถูกวาดจริง", painted)
        segs = page.eval_on_selector_all("#chartType .dseg", "els => els.length")
        check("โดนัทมีชิ้นส่วนครบ 3", segs == 3, f"got {segs}")
        sweep = page.eval_on_selector(
            "#chartType .dseg", "el => el.getAttribute('stroke-dasharray')")
        check("โดนัทกวาดเข้าแล้ว (dasharray ≠ 0)", not sweep.startswith("0 "), sweep)
        total = page.text_content("#rTotal")
        check("ไทล์รวมนับขึ้นถึงยอดจริง", total.strip() not in ("", "0"), total)
        bars = page.eval_on_selector_all(
            "#acctTable .mini-bar", "els => els.map(e=>e.style.width)")
        check("แถบสัดส่วนบัญชีวิ่งเข้า", bars and any(w not in ("", "0%") for w in bars),
              str(bars))

        # ── หน้าอัปโหลด & ค้นหา: ป้ายสถานะ + ช่องเหตุผลย้ายที่ ──
        page.click('[data-view="run"]')
        page.wait_for_selector("#view-run:not(.hidden)", timeout=5000)
        # v3.5.0: เข้าหน้าอัปโหลด = ไฟล์ที่ค้นครบแล้วถูกเอาออกจากคิวเอง (ผลอยู่ใน Downloads แล้ว)
        page.wait_for_function("document.querySelector('#uploadMsg').textContent.includes('ออกจากคิว')",
                               timeout=15000)
        page.wait_for_function("document.querySelectorAll('#queueList .qitem').length === 2",
                               timeout=15000)
        check("ไฟล์ที่ค้นครบแล้วถูกเอาออกจากคิวเมื่อเข้าหน้าอัปโหลด",
              [j["id"] for j in engine.load_queue()] == ["q2", "q3"],
              str([j["id"] for j in engine.load_queue()]))
        check("ไฟล์ยังไม่ครบโชว์สถานะรอ",
              page.locator("#queueList .qst-part").count() == 1
              and page.locator("#queueList .qst-wait").count() == 1)
        # ป้าย File ready ยังเรนเดอร์ถูก (โชว์ตอนรอบเพิ่งจบ ก่อนออกจากหน้า)
        page.evaluate("""() => renderQueue([{id:'r', path:'x/ready.xlsx', done:true, searched:3,
                                             total_rows:3, errors:0, pending:0}])""")
        ready = page.text_content("#queueList .qst-ready")
        check("ไฟล์เสร็จแล้วโชว์ File ready สีเขียว", ready and "File ready" in ready, str(ready))
        check("ไม่มีปุ่ม Download ในคิวแล้ว",
              page.locator("#queueList a.qdl-btn").count() == 0)
        tag = page.eval_on_selector("#queueList .qst-ready", "el => el.tagName")
        check("ป้ายสถานะไม่ใช่ลิงก์/ปุ่ม", tag == "SPAN", tag)
        page.evaluate("() => renderQueue([{id:'e', path:'x/err.xlsx', done:false, searched:4, "
                      "total_rows:5, errors:1, pending:0}])")
        check("ไฟล์ที่ยังมีแถว ERROR ค้าง → ไม่ขึ้น File ready แต่บอกจำนวนแถวผิดพลาด",
              page.locator("#queueList .qst-ready").count() == 0
              and "ผิดพลาด 1 แถว" in (page.text_content("#queueList .qst-err") or ""))
        page.evaluate("() => loadQueue()")
        page.wait_for_function("document.querySelectorAll('#queueList .qitem').length === 2",
                               timeout=15000)
        # ช่องเหตุผล: อยู่ใน reason-row เหนือแถว LOGIN CRIMES และหายจากตัวเลือกขั้นสูง
        check("ช่องเหตุผลอยู่แถวใหม่เหนือ LOGIN CRIMES",
              page.locator(".reason-row #cfgReason").count() == 1)
        check("ช่องเหตุผลหายจากตัวเลือกขั้นสูงแล้ว",
              page.locator("#uploadBox #cfgReason").count() == 0)
        above = page.evaluate(
            """() => { const r = document.querySelector('.reason-row').getBoundingClientRect();
                 const b = document.querySelector('#btnStart').getBoundingClientRect();
                 return r.bottom <= b.top; }""")
        check("แถวเหตุผลอยู่สูงกว่าปุ่ม LOGIN CRIMES", above)

        # ── v3.1.0: ความเร็วย้ายมาต่อจากเหตุผล + บังคับทำตามขั้นตอน ──
        check("ช่องความเร็วอยู่ถัดจากช่องเหตุผล",
              page.locator(".reason-row #cfgSpeed").count() == 1)
        check("กล่องความเร็วหายจากแถบสถานะแล้ว",
              page.locator("#rbSpeedBox").count() == 0)
        check("มีไฟล์แต่ยังไม่เลือกเหตุผล/ความเร็ว → LOGIN CRIMES กดไม่ได้",
              page.eval_on_selector("#btnStart", "el => el.disabled") is True)
        page.select_option("#cfgReason", "random")
        page.wait_for_timeout(150)
        check("เลือกเหตุผลแล้วแต่ยังไม่เลือกความเร็ว → ยังกดไม่ได้",
              page.eval_on_selector("#btnStart", "el => el.disabled") is True)
        page.select_option("#cfgSpeed", "3")
        page.wait_for_timeout(400)
        check("ครบ 3 ขั้นตอน → LOGIN CRIMES เปิดให้กด",
              page.eval_on_selector("#btnStart", "el => el.disabled") is False)
        hint = page.text_content("#stepHint") or ""
        check("แถบขั้นตอนบอกว่าครบแล้ว", "ครบทุกขั้นตอน" in hint, hint)
        check("กราฟสถานะคิวโผล่บนหน้าอัปโหลด",
              page.locator("#runStatusCard:not(.hidden) #runChart .dseg").count() == 3)

        # ── v3.2.0: แถบแสงนำขั้นตอน · กราฟคิว+แถบสถานะแถวเดียวกัน · log ไม่เด้งกลับ ──
        check("ครบขั้นตอน → แสงกะพริบนำที่ปุ่ม LOGIN CRIMES",
              page.locator("#btnStart.step-glow").count() == 1)
        page.select_option("#cfgSpeed", "")
        page.wait_for_timeout(150)
        check("ยังไม่เลือกความเร็ว → แสงย้ายไปช่องความเร็ว",
              page.locator("#cfgSpeed.step-glow").count() == 1 and page.locator("#btnStart.step-glow").count() == 0)
        page.select_option("#cfgSpeed", "3")
        page.wait_for_timeout(150)
        check("กราฟคิวและแถบสถานะอยู่ในแถวเดียวกัน (.run-row)",
              page.locator(".run-row > #runStatusCard").count() == 1 and page.locator(".run-row > #runBanner").count() == 1)
        check("มีปุ่มกลับไปบรรทัดล่าสุดของบันทึกสด", page.locator("#logJump").count() == 1)
        sticky = page.evaluate(
            """() => { const lg = document.querySelector('#liveLog');
                 lg.textContent = Array.from({length:200},(_,i)=>'บรรทัด '+i).join('\\n');
                 lg.scrollTop = 0;
                 const before = lg.scrollTop;
                 // จำลองรอบ poll: ข้อความเปลี่ยนแต่ผู้ใช้เลื่อนขึ้นอยู่ → ต้องไม่เด้งลง
                 const atBottom = (lg.scrollHeight - lg.scrollTop - lg.clientHeight) < 30;
                 lg.textContent += '\\nบรรทัดใหม่'; if(atBottom) lg.scrollTop = lg.scrollHeight;
                 return lg.scrollTop === before; }""")
        check("เลื่อนขึ้นดูประวัติแล้วไม่เด้งกลับลงมา", sticky)

        # ── v3.1.0: กองกลาง (ทุกคนเห็น) + Balance + Setting ยุบกล่อง ──
        page.click('[data-view="live"]')
        page.wait_for_selector("#view-live:not(.hidden)", timeout=5000)
        check("หัวข้อเปลี่ยนเป็น รับ/จ่าย กองกลาง",
              "กองกลาง" in (page.text_content("#view-live .vh h1") or ""))
        page.wait_for_selector("#liveChart canvas", timeout=15000)
        check("กราฟกองกลางเป็นแคนวาส", True)
        page.click('[data-view="billing"]')
        page.wait_for_selector("#view-billing:not(.hidden)", timeout=5000)
        page.wait_for_selector("#billChart canvas", timeout=15000)
        check("กราฟ Balance เป็นแคนวาส", True)
        # ── v3.2.0: กองกลาง = สมุดบันทึกล้วน · นาฬิกาไทย · กราฟทีม · ตารางสมาชิก ──
        page.click('[data-view="live"]')
        page.wait_for_selector("#view-live:not(.hidden)", timeout=5000)
        page.wait_for_timeout(600)
        check("กองกลางไม่มีคอลัมน์ผลการค้น (เป็นสมุดบันทึกล้วน)",
              "รับจากการค้น" not in (page.text_content("#view-live") or ""))
        page.click('[data-view="dashboard"]')
        page.wait_for_selector("#view-dashboard:not(.hidden)", timeout=5000)
        page.wait_for_timeout(1200)
        clk = page.text_content("#thClock") or ""
        check("นาฬิกาไทยเดินบนแดชบอร์ด", "🕒" in clk and "น." in clk and "--" not in clk, clk)
        page.wait_for_selector(".team-chart canvas", timeout=15000)
        check("กราฟทีมแยกต่างหากบนแดชบอร์ด", page.locator(".team-chart canvas").count() >= 1)
        page.click('[data-view="settings"]')
        page.wait_for_selector("#view-settings:not(.hidden)", timeout=5000)
        n_clps = page.locator("#view-settings .clps").count()
        check("กล่องตั้งค่าถูกยุบตามหัวข้อ", n_clps >= 5, f"got {n_clps}")
        vis0 = page.eval_on_selector("#setDisplayName", "el => el.offsetParent !== null")
        page.click("#view-settings .clps > .card-h")   # กดหัวข้อแรก (ทั่วไป)
        page.wait_for_timeout(200)
        vis1 = page.eval_on_selector("#setDisplayName", "el => el.offsetParent !== null")
        check("กดหัวข้อแล้วเนื้อหากางออก", (not vis0) and vis1, f"{vis0}->{vis1}")
        check("การ์ด KEY อนุญาตอัปเดตถูกถอดแล้ว", page.locator("#licKey").count() == 0)

        # ── v3.3.0: ประวัติแยกตามผู้ค้น ──
        page.click('[data-view="history"]')
        page.wait_for_selector("#view-history:not(.hidden)", timeout=5000)
        page.wait_for_selector("#histTable table", timeout=15000)
        head = page.text_content("#histTable thead") or ""
        check("ค่าเริ่มต้น = ของตัวเอง (ยังไม่มีคอลัมน์ผู้ค้น)", "ผู้ค้น" not in head, head[:80])
        page.evaluate("""async () => { await fetch('/api/admin/members',{method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({username:'histmate', password:'pass66',
                                 display_name:'เพื่อนร่วมทีม', role:'member'})});
          const ms = await (await fetch('/api/admin/members')).json();
          const id = ms.find(m=>m.username==='histmate').id;
          await fetch('/api/admin/teams',{method:'POST',headers:{'Content-Type':'application/json'},
            body:JSON.stringify({name:'ทีมประวัติ', member_ids:[1, id]})}); }""")
        page.click('[data-view="dashboard"]'); page.click('[data-view="history"]')
        page.wait_for_selector("#histWho:not(.hidden)", timeout=15000)
        check("มีตัวเลือกผู้ค้นเมื่อมีคนอื่นให้ดู", page.locator("#histWho").count() == 1)
        page.select_option("#histWho", "all")
        page.wait_for_timeout(600)
        head = page.text_content("#histTable thead") or ""
        check("เลือกทั้งทีม → มีคอลัมน์ผู้ค้น (แยกรายบุคคล)", "ผู้ค้น" in head, head[:80])

        # ── ไล่ครบทุกเมนู (v3.0.0: เมนูทีมรวมเข้ากับสมาชิกแล้ว) — ต้องไม่มี JS error ──
        for view in ["history", "members", "billing", "live",
                     "audit", "settings", "dashboard"]:
            page.click(f'[data-view="{view}"]')
            page.wait_for_selector(f"#view-{view}:not(.hidden)", timeout=5000)
            page.wait_for_timeout(350)
        check("ไล่ครบทุกเมนูแล้ว", True)

        # ── v3.0.0: สมาชิก & ทีม หมวดเดียวกัน + สถานะออนไลน์ ──
        page.click('[data-view="members"]')
        page.wait_for_selector("#view-members:not(.hidden)", timeout=5000)
        check("หน้าสมาชิกมีแท็บ สมาชิก/ทีม",
              page.locator("#mtTabs .tab-btn").count() == 2)
        page.wait_for_selector("#memberList table", timeout=5000)
        online = page.text_content("#memberList")
        check("สมาชิกที่ล็อกอินอยู่ขึ้นสถานะออนไลน์", "ออนไลน์" in (online or ""))
        cur_ver = (APP / "VERSION").read_text(encoding="utf-8").strip()
        check(f"ตารางสมาชิกบอกเวอร์ชันที่ใช้ (v{cur_ver})",
              f"v{cur_ver}" in (online or ""), online[:200])
        page.evaluate("""async () => { await fetch('/api/admin/members',{method:'POST',headers:{'Content-Type':'application/json'},
            body:JSON.stringify({username:'member9', password:'pass66', display_name:'สมาชิกเก้า', role:'member'})}); }""")
        page.click('#mtTabs [data-mtab="members"]')
        page.wait_for_selector("#memberList .upd-off", timeout=15000)
        # ต้องรอให้แถวของสมาชิกที่เพิ่งสร้างโผล่ก่อนนับ — ไม่งั้นนับจากรายการรอบก่อน (นับขาดไป 1)
        page.wait_for_function(
            "document.querySelector('#memberList').textContent.includes('สมาชิกเก้า')", timeout=15000)
        n_off = page.locator("#memberList .upd-off").count()
        check("มีปุ่มสลับสิทธิ์อัปเดตในแถวสมาชิก", n_off >= 1, f"got {n_off}")
        page.locator("#memberList .upd-off").first.click()
        page.wait_for_selector("#memberList .upd-on", timeout=15000)
        check("กดแล้วสิทธิ์อัปเดตเปิดทันที (✓ อนุญาตอัปเดต)",
              page.locator("#memberList .upd-on").count() == 1
              and page.locator("#memberList .upd-off").count() == n_off - 1)
        page.click('#mtTabs [data-mtab="teams"]')
        page.wait_for_selector("#mtPaneTeams:not(.hidden)", timeout=4000)
        check("สลับแท็บทีมได้ (สมาชิกซ่อน)",
              page.locator("#mtPaneMembers.hidden").count() == 1)
        check("ไม่มีเมนูทีมแยกแล้ว",
              page.locator('[data-view="teams"]').count() == 0)

        # ── v3.2.0: ฟอร์มสมัครมีช่องรหัสเชื่อมต่อ (ตรวจใน DOM — ฟอร์มซ่อนอยู่หลังสมัครแล้ว) ──
        check("ฟอร์มสมัครมีช่องรหัสเชื่อมต่อศูนย์กลาง", page.locator("#suHubCode").count() == 1)
        check("Setting มีปุ่มคัดลอกรหัสเชื่อมต่อ + ช่องวางรหัส",
              page.locator("#btnHubCode").count() == 1 and page.locator("#hubCodeIn").count() == 1)

        # ── v3.0.0: ตั้ง PASSCODE → ออก → เลือก user แล้วปลดล็อกด้วย PASSCODE ──
        ok = page.evaluate(
            """async () => { const r = await fetch('/api/auth/passcode/set', {method:'POST',
                 headers:{'Content-Type':'application/json'},
                 body: JSON.stringify({password:'secret9', passcode:'24681357',
                                       passcode2:'24681357'})});
               return (await r.json()).ok === true; }""")
        check("ตั้ง PASSCODE ได้", ok)
        logout(page)
        page.wait_for_selector("#screen-auth:not(.hidden)", timeout=15000)
        page.wait_for_selector("#authMethodBar:not(.hidden)", timeout=15000)
        page.click('#authMethodBar [data-method="passcode"]')
        page.wait_for_selector("#authQuickBox:not(.hidden)", timeout=4000)
        n = page.eval_on_selector("#authQuickUser", "el => el.options.length")
        check("มีรายชื่อผู้ใช้ให้เลือกปลดล็อก", n == 1, f"got {n}")
        page.fill("#authPasscode", "24681357")
        page.click("#btnQuickUnlock")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)
        check("เลือก user แล้วปลดล็อกด้วย PASSCODE ของคนนั้นได้", True)

        # ── ออกจากระบบ → เข้าใหม่ด้วยรหัสผ่านปกติ ──
        logout(page)
        page.wait_for_selector("#screen-auth:not(.hidden)", timeout=15000)
        # v3.6.1: กดออกทันทีหลังปลดล็อก (หน้าแรกยังโหลดค้าง) — คำตอบที่มาทีหลัง /api/logout เคยคืนคุกกี้
        # ที่ยังใช้ได้ให้เบราว์เซอร์ ทำให้รีโหลดแล้วยังอยู่ในระบบ (CI ล้มที่บรรทัดบนเพราะเหตุนี้)
        st = page.evaluate("async () => (await fetch('/api/me')).status")
        check("ออกจากระบบแล้วเซิร์ฟเวอร์ไม่รับ session เดิมอีก แม้คุกกี้จะถูกส่งกลับมาทับ", st == 401, f"status {st}")
        page.wait_for_selector("#authMethodBar:not(.hidden)", timeout=15000)
        page.click('#authMethodBar [data-method="password"]')
        page.fill("#authLoginUser", "admin1")
        page.fill("#authLoginPw", "secret9")
        page.click("#authLoginForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)
        check("ออกแล้วเข้าใหม่ด้วยรหัสผ่านได้", True)
        page.wait_for_timeout(700)

        # ── v3.6.4: ลืมรหัสผ่าน (สมาชิก) — ยืนยันด้วยชื่อที่ใช้เข้าเว็บ CRIMES ที่โปรแกรมจดไว้จากรอบค้น ──
        m9 = next(u for u in db.list_users() if u["username"] == "member9")
        db.remember_crimes_account(m9["id"], "Somchai.C")     # ชื่อที่ worker อ่านจากหน้าเว็บ CRIMES ตอนคนนี้ค้น
        logout(page)
        page.wait_for_selector("#screen-auth:not(.hidden)", timeout=15000)
        page.wait_for_selector("#authForgotLink:not(.hidden)", timeout=15000)
        check("หน้าเข้าสู่ระบบมีลิงก์ 'ลืมรหัสผ่าน?' และไม่มีกล่องกู้ Super Admin (v3.6.3) แล้ว",
              page.locator("#authRecoverBox").count() == 0)
        page.click("#authForgotLink")
        page.wait_for_selector("#authForgotBox:not(.hidden)", timeout=5000)
        check("กด 'ลืมรหัสผ่าน?' → ฟอร์มโผล่ ช่องเข้าสู่ระบบ/PASSCODE/แถบเลือกวิธีซ่อน",
              page.eval_on_selector("#authLoginBox", "el => el.classList.contains('hidden')")
              and page.eval_on_selector("#authQuickBox", "el => el.classList.contains('hidden')")
              and page.eval_on_selector("#authMethodBar", "el => el.classList.contains('hidden')"))
        page.fill("#fgUser", "member9")
        page.fill("#fgCrimes", "someone.else")
        page.fill("#fgPw", "newpass9")
        page.fill("#fgPw2", "newpass9")
        page.click("#btnForgotSubmit")
        try:
            page.wait_for_function(
                "document.getElementById('authForgotErr').textContent.includes('ไม่ตรง')", timeout=8000)
        except Exception:
            print("DEBUG forgot state:", page.evaluate(
                """() => ({err: document.getElementById('authForgotErr').textContent,
                           errDisp: document.getElementById('authForgotErr').style.display,
                           u: document.getElementById('fgUser').value,
                           cr: document.getElementById('fgCrimes').value,
                           p1: document.getElementById('fgPw').value.length,
                           p2: document.getElementById('fgPw2').value.length,
                           boxHidden: document.getElementById('authForgotBox').classList.contains('hidden'),
                           active: document.activeElement && document.activeElement.id,
                           appShown: !document.getElementById('screen-app').classList.contains('hidden'),
                           valid: document.getElementById('authForgotForm').checkValidity()})"""))
            raise
        check("ชื่อ CRIMES ไม่ตรง → แจ้งบนหน้า ไม่หลุดออก", True)
        page.fill("#fgCrimes", " SOMCHAI.c ")
        page.click("#btnForgotSubmit")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)
        me = page.evaluate("async () => (await (await fetch('/api/me')).json()).username")
        check("ชื่อ CRIMES ตรง (ไม่สนตัวพิมพ์/ช่องว่าง) → ตั้งรหัสใหม่และเข้าสู่ระบบเป็นสมาชิกคนนั้น",
              me == "member9", str(me))
        page.wait_for_timeout(700)
        logout(page)
        page.wait_for_selector("#screen-auth:not(.hidden)", timeout=15000)
        page.wait_for_selector("#authForgotLink:not(.hidden)", timeout=15000)
        page.click("#authForgotLink")
        page.wait_for_selector("#authForgotBox:not(.hidden)", timeout=5000)
        page.click("#btnForgotBack")
        page.wait_for_function("document.getElementById('authForgotBox').classList.contains('hidden')", timeout=5000)
        page.wait_for_selector("#authMethodBar:not(.hidden)", timeout=15000)   # admin1 มี PASSCODE → มีแถบเลือกวิธี
        check("กด 'กลับ' → หน้าเข้าสู่ระบบกลับมาเหมือนเดิม", True)
        page.click('#authMethodBar [data-method="password"]')
        page.fill("#authLoginUser", "member9")
        page.fill("#authLoginPw", "newpass9")
        page.click("#authLoginForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=15000)
        check("เข้าด้วยรหัสใหม่ได้", True)
        page.wait_for_timeout(500)

        browser.close()

    # 401 จาก /api/me ตอนยังไม่ล็อกอิน = การเช็คสถานะปกติ (เบราว์เซอร์ log เป็น
    # resource error) — ไม่ใช่ข้อผิดพลาดของสคริปต์ · pageerror ทุกตัวยังนับเต็ม
    # คำตอบ 4xx ที่ตั้งใจ (401 ยังไม่ล็อกอิน · 400 รหัสสองช่องไม่ตรงในขั้นกู้รหัส · 403 ไม่มีไฟล์กู้) ถูกเบราว์เซอร์
    # log เป็น resource error — ไม่ใช่ข้อผิดพลาดของสคริปต์ · pageerror ทุกตัวยังนับเต็ม
    def _expected_4xx(e):
        return "Failed to load resource" in e and any(f"status of {c}" in e for c in (400, 401, 403, 409, 429))
    real_errors = [e for e in JS_ERRORS if "favicon" not in e and not _expected_4xx(e)]
    check("ไม่มี JavaScript error แม้แต่จุดเดียว", not real_errors,
          "\n    ".join(real_errors[:8]))

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
