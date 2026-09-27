#!/usr/bin/env python3
"""v3.6.0 — เขียนผลที่คอลัมน์ F เท่านั้น (ไม่เปิดเบราว์เซอร์)

  • detect_columns คืน out_col = F เสมอ แม้ไฟล์มีหัว 'ผล' ที่คอลัมน์อื่น
  • /api/upload ไม่รับ out_col จากฟอร์ม และไม่ใช้ out_col ที่จำไว้จากรุ่นก่อน
  • load_queue บังคับงานเก่าที่จำคอลัมน์อื่นไว้ให้เป็น F · /api/queue รายงาน F
  • ความคืบหน้าของไฟล์อ่านจาก F เท่านั้น (ข้อความในคอลัมน์อื่นไม่นับเป็นผล)
  • หน้าจอไม่มีช่องกรอกคอลัมน์เขียนผลแล้ว และบอกว่าเขียนที่ F
"""
import io
import os
import shutil
import sys
import tempfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v360_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "data")
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "up")
from _app import APP  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)

from backend import server  # noqa: E402
import db  # noqa: E402
import engine  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def make_xlsx(name, headers, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for r, vals in enumerate(rows, 2):
        for c, v in enumerate(vals, 1):
            if v is not None:
                ws.cell(r, c, v)
    p = TEST_DATA / name
    wb.save(p)
    return p


def main():
    print("── ตรวจคอลัมน์ ──")
    check("RESULT_COL คือ F", engine.RESULT_COL == "F")
    p = make_xlsx("a.xlsx", ["ลำดับ", "เลขบัตร", "ผล", "หมายเหตุ"],
                  [[1, 3101501909804, None, None], [2, 3101501909805, None, None]])
    det = engine.detect_columns(p)
    check("มีหัว 'ผล' ที่ C → out_col ยังเป็น F", det and det["out_col"] == "F" and det["id_col"] == "B", str(det))
    check("ไม่มีตัวเดาคอลัมน์ผลเหลืออยู่", not hasattr(engine, "_pick_out_col"))

    print("\n── คิว ──")
    engine.save_queue([{"id": "old", "path": str(p), "user_id": 1, "out_col": "E"},
                       {"id": "new", "path": str(p), "user_id": 1}])
    q = engine.load_queue()
    check("งานเก่าที่จำคอลัมน์ E ไว้ → โหลดมาเป็น F", q[0]["out_col"] == "F" and q[1]["out_col"] == "F")

    print("\n── API ──")
    server.app.config["TESTING"] = True
    c = server.app.test_client()
    c.post("/api/setup", json={"username": "admin1", "password": "secret9",
                               "password2": "secret9", "display_name": "Admin"})
    engine.save_queue([])
    data = {"file": (io.BytesIO(p.read_bytes()), "a.xlsx"), "out_col": "C", "id_col": "B"}
    r = c.post("/api/upload", data=data, content_type="multipart/form-data")
    j = r.get_json()
    check("อัปโหลดส่ง out_col=C มา → งานใช้ F", r.status_code == 200 and j["job"]["out_col"] == "F",
          r.get_data(as_text=True)[:200])
    # รูปแบบที่จำไว้จากรุ่นก่อนบอกว่าผลอยู่ E → ต้องไม่ถูกใช้
    sig, _ = engine.read_header_signature(p)
    db.save_file_format(sig, "B", "E", 2, "")
    engine.save_queue([])
    data = {"file": (io.BytesIO(p.read_bytes()), "a.xlsx")}
    r = c.post("/api/upload", data=data, content_type="multipart/form-data")
    j = r.get_json()
    check("รูปแบบที่จำไว้บอก E → งานยังใช้ F (และจำใหม่เป็น F)",
          j["job"]["out_col"] == "F" and db.get_file_format(sig)["out_col"] == "F", str(j.get("job")))
    q = c.get("/api/queue").get_json()
    check("/api/queue รายงาน out_col F", q and q[0]["out_col"] == "F")

    print("\n── ความคืบหน้าอ่านจาก F ──")
    p2 = make_xlsx("b.xlsx", ["ลำดับ", "เลขบัตร", "ผล", None, None, "F"],
                   [[1, 3101501909804, "คดีที่1 x ปี2560", None, None, None],
                    [2, 3101501909805, None, None, None, " - "]])
    pr = engine.scan_progress(p2, "B", engine.RESULT_COL, 2)
    check("ข้อความที่ C ไม่นับเป็นผล · ผลที่ F นับ", pr["ok"] == 1 and pr["pending"] == 1, str(pr))
    engine.save_queue([{"id": "x", "path": str(p2), "user_id": 1, "out_col": "C"}])
    q = c.get("/api/queue").get_json()
    check("/api/queue ของงานที่จำ C ไว้ ก็นับจาก F", q[0]["searched"] == 1 and q[0]["pending"] == 1, str(q[0]))

    print("\n── หน้าจอ ──")
    html = (APP / "frontend" / "index.html").read_text(encoding="utf-8")
    check("ไม่มีช่องกรอกคอลัมน์เขียนผลแล้ว", 'id="cfgOut"' not in html)
    check("หน้าจอบอกว่าเขียนผลคอลัมน์ F", "เขียนผลคอลัมน์ F" in html)
    check("แถบสถานะโฉมใหม่มีขั้นตอน/วงแหวน/ตัวเลขสด",
          all(k in html for k in ('id="rbSteps"', 'id="rbRingFill"', 'id="rbFound"', 'id="rbEta"', 'data-state')))
    check("ปุ่มควบคุมเดิมยังอยู่ครบ", all(f'id="{b}"' in html for b in (
        "btnConfirmLogin", "btnNextStep", "btnAutoMode", "btnSaveNow", "btnPause", "btnResume",
        "btnRetryErrors", "btnStop", "btnCloseBrowser", "rbDownload", "rbSpeed", "rbIncomplete", "rbRetry")))

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
