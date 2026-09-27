#!/usr/bin/env python3
"""v2.10.0 — ทดสอบ backend: เหตุผลการค้นส่งมากับ /api/run/start ถูกบันทึกลงงานในคิว

กติกา:
  * reason เจาะจง → เขียนทับงานที่ยังไม่ done ของผู้กด (งาน done ห้ามแตะ)
  * reason = "random" หรือว่าง → ไม่แตะงานใด ๆ (คงเหตุผลที่ตั้งตอนอัปโหลดไว้)
  * ความยาวถูกตัดที่ 200 ตัวอักษร
"""
import os
import shutil
import sys
import tempfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v2100_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA)

from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))

from backend import engine, server  # noqa: E402

PASS = FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def main():
    server.app.config["TESTING"] = True
    c = server.app.test_client()

    r = c.post("/api/setup", json={"username": "admin1", "password": "secret9",
                                   "password2": "secret9", "display_name": "Admin"})
    check("สร้าง super admin", r.status_code == 200, r.get_data(as_text=True))

    # เตรียมคิว: งานยังไม่เสร็จ 2 งาน (งานหนึ่งมีเหตุผลเดิม) + งานเสร็จแล้ว 1 งาน
    # v3.5.0: 'เสร็จหรือยัง' อ่านจากไฟล์จริง จึงต้องสร้างไฟล์จริง (a=ยังไม่ค้น, b=ค้นไปบางแถว, c=ครบ)
    import openpyxl
    up = TEST_DATA / "up"
    up.mkdir()
    for name, filled in (("a.xlsx", 0), ("b.xlsx", 1), ("c.xlsx", 3)):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(1, 2, "เลขบัตร")
        for i in range(3):
            ws.cell(2 + i, 2, 3101501909800 + i)
            if i < filled:
                ws.cell(2 + i, 6, " - ")
        wb.save(up / name)
    engine.save_queue([
        {"id": "j1", "path": str(up / "a.xlsx"), "user_id": 1},
        {"id": "j2", "path": str(up / "b.xlsx"), "user_id": 1,
         "reason": "เหตุผลตอนอัปโหลด"},
        {"id": "j3", "path": str(up / "c.xlsx"), "user_id": 1,
         "reason": "ของงานที่จบแล้ว"},
    ])

    # ไม่ให้รอบทำงานจริงถูกปล่อย — จับพารามิเตอร์แทน
    calls = []
    server.manager.is_busy = lambda: False
    server.manager.start = lambda jobs, account, **kw: (calls.append(jobs), (True, "ok"))[1]

    # 1) เหตุผลเจาะจง → เขียนทับงานที่ยังไม่เสร็จ, งาน done ไม่ถูกแตะ
    r = c.post("/api/run/start", json={"account": "", "reason": "ตรวจสอบประวัติผู้สมัครงาน"})
    check("start ตอบ ok", r.status_code == 200 and r.get_json().get("ok") is True,
          r.get_data(as_text=True))
    q = {j["id"]: j for j in engine.load_queue()}
    check("งานใหม่ได้เหตุผลใหม่", q["j1"].get("reason") == "ตรวจสอบประวัติผู้สมัครงาน")
    check("งานที่มีเหตุผลเดิมถูกเขียนทับ", q["j2"].get("reason") == "ตรวจสอบประวัติผู้สมัครงาน")
    check("งานที่จบแล้วไม่ถูกแตะ", q["j3"].get("reason") == "ของงานที่จบแล้ว")
    check("manager.start ได้รับงานพร้อมเหตุผลใหม่",
          calls and all(j.get("reason") == "ตรวจสอบประวัติผู้สมัครงาน"
                        for j in calls[0] if not j.get("done")))
    check("v3.5.0: ไฟล์ที่ค้นครบแล้วไม่ถูกส่งไปรันซ้ำ",
          calls and sorted(j["id"] for j in calls[0]) == ["j1", "j2"],
          str([j["id"] for j in calls[0]]) if calls else "")

    # 2) random → ไม่แตะเหตุผลที่มีอยู่
    r = c.post("/api/run/start", json={"account": "", "reason": "random"})
    check("start(random) ตอบ ok", r.status_code == 200)
    q = {j["id"]: j for j in engine.load_queue()}
    check("random ไม่เขียนทับ", q["j2"].get("reason") == "ตรวจสอบประวัติผู้สมัครงาน")

    # 3) ไม่ส่ง reason เลย (ไคลเอนต์เก่า) → ทำงานได้ ไม่แตะคิว
    r = c.post("/api/run/start", json={"account": ""})
    check("ไคลเอนต์เก่าไม่ส่ง reason ยังใช้ได้", r.status_code == 200)

    # 4) เหตุผลยาวเกิน → ถูกตัดที่ 200 ตัวอักษร
    long_reason = "ย" * 500
    c.post("/api/run/start", json={"account": "", "reason": long_reason})
    q = {j["id"]: j for j in engine.load_queue()}
    check("เหตุผลยาวถูกตัดที่ 200", len(q["j1"].get("reason", "")) == 200)

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
