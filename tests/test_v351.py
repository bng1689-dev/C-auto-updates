#!/usr/bin/env python3
"""v3.5.1 — พบคดีแต่ไม่มีข้อหา → ทาเหลืองทั้งแถว (ไม่เปิดเบราว์เซอร์)

  • charge_missing: ว่าง/ขีด/คำว่าไม่ระบุ นับเป็นไม่มีข้อหา
  • review_kind: ไม่มีข้อหาชนะไม่มีปี · review_kind_from_text อ่านจากข้อความผลได้ (กู้จากประวัติ)
  • apply_row_review: เหลืองทั้งแถว A..F รวมเซลล์เลขบัตร · ไม่ทับมาร์คเติมศูนย์/เลขไม่ตรง
    · ไม่มีปี = แดงอ่อน ไม่แตะเซลล์เลขบัตร · ล้างเฉพาะสีของระบบ ไม่ลบสีของผู้ใช้
  • สลับสถานะข้ามรอบ: แถวเคยไม่มีข้อหา → รอบใหม่ครบ = สีหาย / เคยแดงอ่อน → ไม่มีข้อหา = เหลือง
"""
import os
import shutil
import sys
import tempfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_test_v351_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "data")
from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))

import engine  # noqa: E402

PASS = FAIL = 0
YEL, RED, LRED, NONE = "FFE699", "FF6B6B", "FFC7CE", ""


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


def fill(ws, r, c):
    rgb = ws.cell(r, c).fill.start_color.rgb if ws.cell(r, c).fill.fill_type else ""
    return (rgb or "")[-6:] if isinstance(rgb, str) else ""


def row_fills(ws, r, upto=6):
    return [fill(ws, r, c) for c in range(1, upto + 1)]


def main():
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    print("── ตัดสินว่า 'ไม่มีข้อหา' ──")
    for v, want in (("", True), (None, True), ("-", True), ("  —  ", True), ("ไม่ระบุ", True),
                    ("(ไม่ระบุข้อหา)", True), ("ลักทรัพย์", False), ("ลักทรัพย์/รับของโจร", False)):
        check(f"charge_missing({v!r}) = {want}", engine.charge_missing(v) == want)
    kept = [{"charge": "ลักทรัพย์", "year": "2560"}, {"charge": "", "year": "2561"}]
    check("คดีใดคดีหนึ่งไม่มีข้อหา → 'charge'", engine.review_kind(kept) == "charge")
    check("มีข้อหาแต่ไม่มีปี → 'year'", engine.review_kind([{"charge": "x", "year": ""}]) == "year")
    check("ไม่มีทั้งข้อหาและปี → 'charge' (สำคัญกว่า)",
          engine.review_kind([{"charge": "", "year": ""}]) == "charge")
    check("ครบ → None", engine.review_kind([{"charge": "x", "year": "2560"}]) is None)
    check("ไม่ใส่ปี: ไม่มีปีไม่นับ", engine.review_kind([{"charge": "x", "year": ""}], no_year=True) is None)
    text, review = engine.format_cases(kept)
    check("ข้อความผลเขียน '(ไม่ระบุข้อหา)' แทนข้อหาว่าง",
          text == "คดีที่1 ลักทรัพย์ ปี2560/ คดีที่2 (ไม่ระบุข้อหา) ปี2561" and review, text)
    text2, _ = engine.format_cases([{"charge": "-", "year": "2560"}])
    check("ข้อหาเป็นขีด → เขียน '(ไม่ระบุข้อหา)' ด้วย", text2 == "คดีที่1 (ไม่ระบุข้อหา) ปี2560", text2)
    check("อ่านชนิดจากข้อความผล (กู้จากประวัติ): มี (ไม่ระบุข้อหา) → charge",
          engine.review_kind_from_text(text, incomplete=1) == "charge")
    check("ข้อความครบแต่ธง incomplete → year",
          engine.review_kind_from_text("คดีที่1 x ปี", incomplete=1) == "year")
    check("ข้อความครบ ไม่มีธง → None", engine.review_kind_from_text("คดีที่1 x ปี2560", 0) is None)

    print("\n── ทาสีในไฟล์ ──")
    wb = Workbook()
    ws = wb.active
    for r in range(2, 8):
        for c in range(1, 7):
            ws.cell(r, c, f"r{r}c{c}")
    ID, OUT = 2, 5
    engine.apply_row_review(ws, 2, OUT, ID, "charge")
    check("ไม่มีข้อหา → เหลืองทั้งแถว A..F รวมเซลล์เลขบัตร", row_fills(ws, 2) == [YEL] * 6, row_fills(ws, 2))
    ws.cell(3, ID).fill = engine.YELLOW              # เคยเติมศูนย์
    engine.apply_row_review(ws, 3, OUT, ID, "charge")
    check("ไม่ทับมาร์คเติมศูนย์ (FFFF00) ของเซลล์เลขบัตร",
          row_fills(ws, 3) == [YEL, "FFFF00", YEL, YEL, YEL, YEL], row_fills(ws, 3))
    ws.cell(4, ID).fill = engine.RED                 # เลขไม่ตรง
    engine.apply_row_review(ws, 4, OUT, ID, "charge")
    check("ไม่ทับมาร์คเลขไม่ตรง (FF6B6B) ของเซลล์เลขบัตร", fill(ws, 4, ID) == RED and fill(ws, 4, 1) == YEL)
    engine.apply_row_review(ws, 5, OUT, ID, "year")
    check("ไม่มีปี → แดงอ่อนทั้งแถว ยกเว้นเซลล์เลขบัตร",
          row_fills(ws, 5) == [LRED, NONE, LRED, LRED, LRED, LRED], row_fills(ws, 5))
    ws.cell(6, 1).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    engine.apply_row_review(ws, 6, OUT, ID, None)
    check("ล้างมาร์ค: ไม่ลบสีที่ผู้ใช้ทาเอง", fill(ws, 6, 1) == "C6EFCE")
    engine.apply_row_review(ws, 7, OUT, ID, "charge")
    engine.apply_row_review(ws, 7, OUT, ID, None)
    check("รอบใหม่ข้อมูลครบ → เหลืองของระบบถูกล้างทั้งแถว", row_fills(ws, 7) == [NONE] * 6, row_fills(ws, 7))
    engine.apply_row_review(ws, 5, OUT, ID, "charge")
    check("เคยแดงอ่อน → รอบใหม่ไม่มีข้อหา = เหลืองทั้งแถว (แดงอ่อนหายหมด)",
          row_fills(ws, 5) == [YEL] * 6, row_fills(ws, 5))
    engine.apply_row_review(ws, 2, OUT, ID, "year")
    check("เคยเหลือง → รอบใหม่มีข้อหาแต่ไม่มีปี = แดงอ่อน เซลล์เลขบัตรถูกล้าง",
          row_fills(ws, 2) == [LRED, NONE, LRED, LRED, LRED, LRED], row_fills(ws, 2))
    engine.apply_row_review(ws, 2, 9, ID, "charge")
    check("คอลัมน์ผลไกลกว่า F → ทาถึงคอลัมน์ผล", fill(ws, 2, 9) == YEL and fill(ws, 2, 7) == YEL)
    ws.merge_cells("A8:C8")
    engine.apply_row_review(ws, 8, OUT, ID, "charge")
    check("ช่วง merge ทาที่เซลล์มุมบนซ้าย", fill(ws, 8, 1) == YEL)
    # mark_review_row เดิมยังล้างเหลืองของแถวได้ (โค้ดเก่าที่ยังเรียกอยู่ ถ้ามี)
    engine.mark_review_row(ws, 8, OUT, on=False, skip_cols=(ID,))
    check("mark_review_row(on=False) ล้างเหลืองของระบบได้ด้วย", fill(ws, 8, 1) == NONE)

    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
finally:
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
