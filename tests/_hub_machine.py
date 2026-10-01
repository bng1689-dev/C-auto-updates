"""เครื่องจำลอง 'อีกเครื่องหนึ่ง' สำหรับ test_v370 — โปรเซสแยก ฐานข้อมูลของตัวเอง คุยกับศูนย์กลางจำลอง (Node) ใบเดียวกัน

    python _hub_machine.py <data_dir> <state_file> <token> [<script_admin_token>]
        (คำสั่งเป็น JSON list ทาง stdin · ตอบ JSON ทาง stdout)
    script_admin_token = ADMIN_TOKEN ที่ฝังในสคริปต์จำลอง (ไม่ใช่ของเครื่องนี้ — เครื่องนี้ใส่ผ่าน op set_admin_token)

ไม่ใช่ชุดทดสอบ — run_all ไม่รันไฟล์นี้ตรง ๆ"""
import json
import os
import subprocess
import sys
from pathlib import Path

data_dir, state_file, token = sys.argv[1:4]
script_admin = sys.argv[4] if len(sys.argv) > 4 else ""
os.environ["CRIMES_DATA_DIR"] = data_dir
os.environ["CRIMES_UPLOAD_DIR"] = str(Path(data_dir) / "up")
os.environ["CRIMES_KEEP_SEED"] = "1"     # เครื่องจำลองอาจกำลังทดสอบไฟล์ seed ที่เพิ่งถูกวาง — _app ห้ามเก็บกวาดทิ้ง
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from _app import APP  # noqa: E402,F401
import hub  # noqa: E402


def fake_post(url, payload, tok, admin_token=None, timeout=None):
    body = hub.encode_payload(payload).decode("ascii")   # ก้อนจริงเป็น ASCII ล้วน (v3.10.2 — ดู hub.encode_payload)
    param = {"sign": hub.sign(body.encode("utf-8"), tok)}
    if admin_token:
        param["asign"] = hub.sign(body.encode("utf-8"), admin_token)
    req = {"method": "POST", "body": body, "parameter": param}
    r = subprocess.run(["node", str(HERE / "gas_harness.js"), state_file, token, script_admin],
                       input=json.dumps(req), capture_output=True, text=True, timeout=60)
    if r.returncode:
        raise RuntimeError("harness: " + r.stderr[-300:])
    return json.loads(r.stdout)


# ต้องสวมก่อน import server — v3.10.0 server ดึงไดเรกทอรีกลางตั้งแต่ตอน import (_boot_seed) ห้ามให้แตะเน็ตจริง
hub._post = fake_post
from backend import server  # noqa: E402
import db  # noqa: E402

server._hub_kick = lambda *a, **k: None          # ซิงก์เมื่อสั่งเท่านั้น (ให้ผลทดสอบแน่นอน)
server.app.config["TESTING"] = True
c = server.app.test_client()


def uid_of(username):
    u = db.get_user_by_username(username)
    return u["id"] if u else None


out = {}
for cmd in json.loads(sys.stdin.read() or "[]"):
    op = cmd["op"]
    key = cmd.get("as", op)
    if op == "join":
        r = c.post("/api/setup/join", json={"hub_code": cmd["code"]})
        out[key] = [r.status_code, r.get_json()]
    elif op == "login":
        r = c.post("/api/login", json={"username": cmd["u"], "password": cmd["p"]})
        out[key] = [r.status_code, r.get_json()]
    elif op == "logout":
        out[key] = c.post("/api/logout").status_code
    elif op == "me":
        r = c.get("/api/me")
        out[key] = [r.status_code, (r.get_json() or {}).get("username")]
    elif op == "password":                       # เปลี่ยนรหัสผ่านของคนที่ล็อกอินอยู่ (ต้องรู้รหัสเดิม)
        r = c.post("/api/password", json={"current": cmd["cur"], "new": cmd["new"]})
        out[key] = [r.status_code, r.get_json()]
    elif op == "create":
        r = c.post("/api/admin/members", json={"username": cmd["u"], "password": cmd["p"],
                                               "display_name": cmd.get("name", ""), "role": cmd.get("role", "member")})
        out[key] = [r.status_code, r.get_json()]
    elif op == "update":
        r = c.put(f"/api/admin/members/{uid_of(cmd['u'])}", json=cmd["fields"])
        out[key] = [r.status_code, r.get_json()]
    elif op == "delete":
        r = c.delete(f"/api/admin/members/{uid_of(cmd['u'])}")
        out[key] = [r.status_code, r.get_json()]
    elif op == "set_admin_token":                # ใส่/ถอด 'รหัสผู้ดูแลศูนย์กลาง' ของเครื่องนี้
        server.auth.update_config(hub_admin_token=str(cmd.get("token") or ""))
        out[key] = True
    elif op == "touch":                          # จำลองนาฬิกาเครื่องเพี้ยน — ตั้ง updated_at ของคนนี้เป็นค่าที่กำหนด
        with db.get_conn() as conn:
            conn.execute("UPDATE users SET updated_at=?, sync_dirty=1 WHERE username=?", (cmd["at"], cmd["u"]))
        out[key] = True
    elif op == "sync":
        res = server._members_sync_now("test")
        out[key] = {"ok": bool(res.get("ok")), "error": res.get("error", ""), "summary": res.get("summary"),
                    "rejected": res.get("rejected"), "admin": res.get("admin")}
    elif op == "pending":
        out[key] = db.count_pending_members()
    elif op == "export":
        out[key] = db.export_members()
    elif op == "user":
        u = db.get_user_by_username(cmd["u"])
        out[key] = ({k: u.get(k) for k in ("display_name", "role", "active", "hub_rev", "sync_dirty", "updated_at")}
                    if u else None)
    elif op == "members":
        out[key] = [{"username": u["username"], "display_name": u["display_name"], "role": u["role"],
                     "active": u["active"], "pending": u.get("pending", 0)} for u in db.list_users(include_pending=True)]
    elif op == "state":
        out[key] = c.get("/api/hub/members/state").get_json()
    elif op == "has_passcode":
        out[key] = db.has_passcode(uid_of(cmd["u"]))
    # ---- v3.9.0: สมุดกองกลางกลาง ----
    elif op == "ledger_add":                     # บันทึกรายการให้เจ้าของ u (คนที่ล็อกอินอยู่ต้องเป็น Super Admin)
        r = c.post("/api/admin/ledger", json={"user_id": uid_of(cmd["u"]), "kind": cmd["kind"], "amount": cmd["amount"],
                                              "note": cmd.get("note", ""), "ym": cmd.get("ym", "")})
        out[key] = [r.status_code, r.get_json()]
    elif op == "ledger_del":                     # ลบด้วย gid (id ในเครื่องต่างกันแต่ละเครื่อง)
        with db.get_conn() as conn:
            row = conn.execute("SELECT id FROM ledger WHERE gid=?", (cmd["gid"],)).fetchone()
        r = c.delete(f"/api/admin/ledger/{row['id'] if row else 0}")
        out[key] = [r.status_code, r.get_json()]
    elif op == "ledger_sync":
        res = server._ledger_sync_now("test")
        out[key] = {"ok": bool(res.get("ok")), "error": res.get("error", ""), "summary": res.get("summary"),
                    "rejected": res.get("rejected"), "admin": res.get("admin"), "seq": res.get("seq")}
    elif op == "ledger":                         # ทุกแถวในเครื่อง (รวมที่ลบแล้ว) พร้อมธงซิงก์
        with db.get_conn() as conn:
            out[key] = [dict(r) for r in conn.execute(
                "SELECT gid, owner, kind, amount, note, deleted, hub_rev, sync_dirty, user_id, origin FROM ledger ORDER BY id").fetchall()]
    elif op == "ledger_pending":
        out[key] = db.count_pending_ledger()
    elif op == "board":                          # หน้ากองกลางในสายตาคนที่ล็อกอินอยู่
        r = c.get("/api/live/board" + (f"?ym={cmd['ym']}" if cmd.get("ym") else ""))
        out[key] = [r.status_code, r.get_json()]
    elif op == "ledger_state":
        out[key] = dict(server._ledger_last)
    elif op == "teams":                          # ทีมในเครื่องนี้: ชื่อ · อัตรา · สมาชิก (ชื่อผู้ใช้) · สร้างตามกลางไหม
        with db.get_conn() as conn:
            synced = {r["id"]: r["synced"] for r in conn.execute("SELECT id, synced FROM teams").fetchall()}
        out[key] = [{"name": t["name"], "rate": t["rate_per_name"], "synced": synced.get(t["id"], 0),
                     "members": sorted(m["username"] for m in t["members"])} for t in db.list_teams()]
    # ---- v3.10.0: สมัครใช้งานเอง ----
    elif op == "register":
        r = c.post("/api/register", json={"username": cmd["u"], "password": cmd["p"], "password2": cmd.get("p2", cmd["p"]),
                                          "display_name": cmd.get("name", ""), "hub_code": cmd.get("code", "")})
        out[key] = [r.status_code, r.get_json()]
    elif op == "auth_state":
        out[key] = c.get("/api/auth/state").get_json()
    elif op == "clear_throttle":                 # ล้างตัวนับหน่วงสมัคร/ล็อกอิน (ชุดทดสอบยิงติดกันหลายครั้ง)
        with db.get_conn() as conn:
            conn.execute("DELETE FROM auth_throttle")
        out[key] = True
    elif op == "login_full":                     # เหมือน login แต่คืน JSON เต็ม (ดูข้อความ 'รออนุมัติ')
        r = c.post("/api/login", json={"username": cmd["u"], "password": cmd["p"]})
        out[key] = [r.status_code, r.get_json()]
    # ---- v3.11.0: ยืนยันผลรายไฟล์ + ส่งขึ้น Google Drive ----
    elif op == "make_file":                      # สร้างไฟล์งาน .xlsx (เลขบัตรถูกต้อง n แถว) — คืน path
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="ลำดับ")
        ws.cell(row=1, column=2, value="เลขบัตรประชาชน")
        results = cmd.get("results") or []
        for i in range(int(cmd["n"])):
            ws.cell(row=2 + i, column=1, value=i + 1)
            base = [int(d) for d in f"3{int(cmd.get('seed', 0)) % 10}{i:010d}"]   # 12 หลักแรก (ไม่ซ้ำรายแถว)
            check = (11 - sum(d * (13 - j) for j, d in enumerate(base)) % 11) % 10
            ws.cell(row=2 + i, column=2, value="".join(map(str, base)) + str(check))
            if i < len(results) and results[i]:
                ws.cell(row=2 + i, column=6, value=results[i])
        src = Path(data_dir) / "src"
        src.mkdir(parents=True, exist_ok=True)
        fp = src / cmd["name"]
        wb.save(fp)
        out[key] = str(fp)
    elif op == "upload_api":                     # อัปโหลดเข้าโปรแกรมจริง (multipart) — เข้าคิวเหมือนผู้ใช้ลากไฟล์
        p = Path(cmd["path"])
        with open(p, "rb") as fh:
            r = c.post("/api/upload", data={"file": (fh, p.name)}, content_type="multipart/form-data")
        out[key] = [r.status_code, r.get_json()]
    elif op == "fill_results":                   # จำลองผลที่ worker เขียนลงคอลัมน์ F (ไฟล์ในคิว)
        import openpyxl
        fp = Path(cmd["path"])
        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        for i, v in enumerate(cmd["values"]):
            if v:
                ws.cell(row=2 + i, column=6, value=v)
        wb.save(fp)
        out[key] = True
    elif op == "http_delete":                    # DELETE ตาม path (op 'delete' เดิมจองไว้ให้ลบสมาชิก)
        r = c.delete(cmd["path"])
        out[key] = [r.status_code, r.get_json()]
    elif op == "drive_direct":                   # เรียก hub.drive_config/upload_file ตรง ๆ (ทดสอบชั้น hub)
        cfg = server.auth.load_config()
        if cmd.get("set") is not None:
            res = hub.drive_config(cfg.get("hub_url"), cfg.get("hub_token"),
                                   admin_token=cfg.get("hub_admin_token") or None, set_cfg=cmd["set"])
        else:
            res = hub.drive_config(cfg.get("hub_url"), cfg.get("hub_token"),
                                   admin_token=cfg.get("hub_admin_token") or None)
        out[key] = res
    elif op in ("get", "post", "put", "delete"):  # เรียก API ใดก็ได้ (ใช้เตรียมสถานการณ์)
        r = getattr(c, op)(cmd["path"], json=cmd.get("json")) if op != "get" else c.get(cmd["path"])
        out[key] = [r.status_code, r.get_json()]
    else:
        out[key] = f"unknown op {op}"
print(json.dumps(out, ensure_ascii=False))
