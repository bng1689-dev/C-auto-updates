#!/usr/bin/env python3
"""ทดสอบทั้งระบบแบบ end-to-end (v3.5.0) — โปรแกรมจริง + worker จริง + Chromium จริง 2 ตัว

  ตัวที่ 1 = หน้าจอ CRIMES AUTO (ผู้ใช้กดจริง)   ตัวที่ 2 = เว็บ CRIMES จำลอง (worker ขับเอง)

ไล่เส้นทางจริงทั้งหมด: อัปโหลดไฟล์ → เลือกเหตุผล/ความเร็ว → LOGIN CRIMES → รอเข้าสู่ระบบ
→ ตรวจพบว่าเข้าสู่ระบบแล้วเริ่มเอง → ปิดกล่อง 'มีอะไรใหม่' → ค้นทีละแถว → แถว ERROR ถูกย้อนกลับ
ไปค้นใหม่ → ไฟล์ครบ → ออกจากคิวเมื่อกลับเข้าหน้าอัปโหลด + กรณีผู้ใช้ปิดหน้าต่าง Chrome เอง

ครอบคลุม 4 ปัญหาที่ผู้ใช้รายงาน:
  (1) ไม่มีปุ่มกดเริ่ม  (2) ไฟล์ที่เสร็จแล้วค้างในคิว  (3) แถว ERROR ไม่ถูกย้อนกลับไปแก้
  (4) เลข 12 หลักไม่ถูกเติม 0 ข้างหน้า
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="crimes_e2e_"))
os.environ["CRIMES_DATA_DIR"] = str(TEST_DATA / "data")
os.environ["CRIMES_UPLOAD_DIR"] = str(TEST_DATA / "uploads")
HERE = Path(__file__).parent
from _app import APP, CHROME  # noqa: E402  (โค้ดแอปจาก build/ = แพ็กเกจรุ่นล่าสุด)
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(APP))
sys.path.insert(0, str(APP / "backend"))

from backend import server  # noqa: E402
import engine  # noqa: E402   (โมดูลเดียวกับที่ worker/server ใช้จริง)
import worker  # noqa: E402
import mock_crimes  # noqa: E402

APORT, MPORT = 8801, 8802
PROFILE = TEST_DATA / "chrome_profile"
DOWNLOADS = TEST_DATA / "Downloads"
DOWNLOADS.mkdir()
PASS = FAIL = 0
JS_ERRORS = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


# ---------------- ต่อ engine เข้ากับเว็บจำลอง ----------------
def fake_launch(p, headless=False):
    return p.chromium.launch_persistent_context(
        str(PROFILE), executable_path=CHROME, headless=True,
        args=["--no-sandbox", "--disable-background-networking"])


engine.launch_browser = fake_launch
engine.kill_profile_browsers = lambda: 0
engine.probe_online = lambda timeout=4: True
engine.SEARCH_URL = f"http://127.0.0.1:{MPORT}/bdasearch/#/bda/search/criteria/person"
engine.SEARCH_HOST = f"127.0.0.1:{MPORT}"
engine.SPEED_FACTORS = {k: 0.05 for k in range(1, 6)}
worker.RunManager.RETRY_BACKOFF = (0, 0)
worker.RunManager.AUTOSTART_SEC = 2
worker.downloads_dir = lambda: DOWNLOADS


def kill_crimes_browser():
    subprocess.run(["pkill", "-f", str(PROFILE)], capture_output=True)


# ---------------- ไฟล์ทดสอบ ----------------
ROWS = [  # (แถว, ค่าเลขบัตร, ช่องผลเดิม)
    (2, 3101501909804, None),                 # 13 หลักปกติ → พบ 1 คดี (อีกคดีเป็นผู้เสียหาย)
    (3, 123456789012, None),                  # 12 หลัก → เติม 0 → พบ (ตารางเก่าค้าง: กันอ่านผลคนก่อน)
    (4, "1-1002-00345-67-8", None),           # มีขีด → ไม่พบ
    (5, "๓๑๐๑๕๐๑๙๐๙๘๐๕", None),               # เลขไทย + เว็บตอบ 500 ครั้งแรก → ERROR → ย้อนกลับไปค้นใหม่
    (6, "123456789", None),                   # 9 หลัก → ใช้ไม่ได้ (หมายเหตุ ไม่ค้น)
    (7, 3101501909806, "ERROR!: หน้าค้นหา CRIMES เปลี่ยนไป — ไม่พบช่องกรอกเลขบัตร (#inputPid) อาจต้องอัปเดต selector"),
    (8, 3101501909807, "ERROR#3: Timeout 30000ms exceeded."),
    (9, 987654321098, "คดีที่1 เดิม ปี2550"),   # มีผลแล้ว แต่ 12 หลัก → ต้องเติม 0 ด้วย
    (10, "FORMULA", None),                    # =C10 (ค่าที่คำนวณไว้ 3101501909808)
    (11, "FLOATSCI", None),                   # ทศนิยมแบบ E+ ในไฟล์ (1.23456789013E+11) + ป๊อบอัพค้างครั้งแรก
    (12, 3101501909809, None),
]


def make_file(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["ลำดับ", "เลขบัตรประชาชน", "ช่วย", "หมายเหตุ", "ผล"], 1):
        ws.cell(1, c, h)
    for r, v, res in ROWS:
        ws.cell(r, 1, r - 1)
        if v == "FORMULA":
            ws.cell(r, 2, 111)        # ตัวแทน — แทนด้วย XML สูตรทีหลัง
            ws.cell(r, 3, 3101501909808)
        elif v == "FLOATSCI":
            ws.cell(r, 2, 222)        # ตัวแทน
        else:
            ws.cell(r, 2, v)
        if res:
            ws.cell(r, 5, res)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    # แก้ XML ตรง ๆ ให้เหมือนไฟล์จริงจากโปรแกรมอื่น: สูตรที่มีค่าคำนวณไว้ + ตัวเลขทศนิยมแบบ E+
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                s = data.decode()
                s, n1 = re.subn(r'<c r="B10"[^>]*>.*?</c>',
                                '<c r="B10"><f>C10</f><v>3101501909808</v></c>', s)
                s, n2 = re.subn(r'<c r="B11"[^>]*>.*?</c>',
                                '<c r="B11"><v>1.23456789013E+11</v></c>', s)
                assert n1 == 1 and n2 == 1, (n1, n2)
                data = s.encode()
            zout.writestr(item, data)
    tmp.unlink()


def status():
    return server.manager.status()


def wait_state(pred, timeout=90, what=""):
    end = time.time() + timeout
    while time.time() < end:
        s = status()
        if pred(s):
            return s
        time.sleep(0.3)
    raise AssertionError(f"หมดเวลารอ {what}: state={status()['state']} msg={status()['message']}")


def main():
    mock_crimes.serve(MPORT)
    S = mock_crimes.STATE
    S.update(logged_in=False, announce=True, delay=0.8,
             fail_http={"3101501909805": 1}, stuck=["0123456789013"],
             cases={
                 "3101501909804": [
                     {"caseNo": "11/2560", "year": "2560", "charge": "ลักทรัพย์", "status": "ผู้ต้องหา"},
                     {"caseNo": "12/2561", "year": "2561", "charge": "ฉ้อโกง", "status": "ผู้เสียหาย"}],
                 "0123456789012": [
                     {"caseNo": "5/2565", "year": "2565", "charge": "ยาเสพติด", "status": "ผู้ต้องหา"}],
                 "3101501909806": [
                     {"caseNo": "7/2561", "year": "2561", "charge": "ฉ้อโกง", "status": "ผู้ต้องหา"}],
                 "0123456789013": [
                     {"caseNo": "9/2563", "year": "2563", "charge": "ปลอมแปลงเอกสาร", "status": "ผู้ต้องหา"}],
                 # v3.5.1: มีข้อหาแต่ไม่มีปี (เลขคดีไม่มี /ปี ให้ดึง) → แดงอ่อนทั้งแถว ยกเว้นเซลล์เลขบัตร
                 "3101501909807": [
                     {"caseNo": "8/62", "year": "", "charge": "ทำร้ายร่างกาย", "status": "ผู้ต้องหา"}],
                 # v3.5.1: พบคดีแต่ไม่มีข้อหา → เหลืองทั้งแถว A..F
                 "3101501909809": [
                     {"caseNo": "3/2559", "year": "2559", "charge": "", "status": "ผู้ต้องหา"}],
             })

    server.app.config["TESTING"] = False
    import logging
    logging.getLogger("werkzeug").setLevel(logging.ERROR)
    threading.Thread(target=lambda: server.app.run(port=APORT, threaded=True, use_reloader=False),
                     daemon=True).start()
    time.sleep(1.2)

    src = TEST_DATA / "งานทดสอบ.xlsx"
    make_file(src)

    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path=CHROME, headless=True,
                               args=["--disable-background-networking"])
        page = b.new_page(viewport={"width": 1280, "height": 860})
        page.on("pageerror", lambda e: JS_ERRORS.append(f"pageerror: {e}"))
        page.on("dialog", lambda d: (JS_ERRORS.append("dialog: " + d.message)
                                     if "เริ่มไม่สำเร็จ" in d.message else None, d.accept()))
        page.goto(f"http://127.0.0.1:{APORT}", wait_until="domcontentloaded")
        page.wait_for_selector("#authSetupBox:not(.hidden)", timeout=8000)
        page.fill("#suUser", "admin1")
        page.fill("#suName", "แอดมิน")
        page.fill("#suPw", "secret9")
        page.fill("#suPw2", "secret9")
        page.click("#authSetupForm button[type=submit]")
        page.wait_for_selector("#screen-app:not(.hidden)", timeout=8000)

        print("\n── อัปโหลด + ขั้นตอน ①②③ ──")
        page.click('[data-view="run"]')
        page.wait_for_selector("#view-run:not(.hidden)")
        page.set_input_files("#fileInput", str(src))
        page.wait_for_function("document.querySelector('#uploadMsg').textContent.includes('เพิ่มแล้ว')",
                               timeout=10000)
        msg = page.inner_text("#uploadMsg")
        check("อัปโหลดแล้วบอกจำนวนแถวถูก (11 แถว)", "11 แถว" in msg, msg)
        check("บอกล่วงหน้าว่าจะเติม 0 ให้เลข 12 หลัก (3 แถว: int, มีผลแล้ว, ทศนิยม E+)",
              "เลข 12 หลัก 3 แถว" in msg, msg)
        check("บอกล่วงหน้าว่ามีเลขใช้ไม่ได้ 1 แถว", "ใช้ไม่ได้ 1 แถว" in msg, msg)
        page.wait_for_selector("#queueList .qitem")
        page.select_option("#cfgReason", "random")
        page.select_option("#cfgSpeed", "5")
        page.wait_for_function("!document.querySelector('#btnStart').disabled", timeout=5000)

        print("\n── (1) ปุ่มเริ่ม: LOGIN CRIMES → รอเข้าสู่ระบบ → ตรวจพบแล้วเริ่มเอง ──")
        page.click("#btnStart")
        wait_state(lambda s: s["state"] == "awaiting_login", 60, "awaiting_login")
        page.wait_for_selector("#btnConfirmLogin:not(.hidden)", timeout=8000)
        vis = page.evaluate("""() => { const b = document.querySelector('#btnConfirmLogin');
            const r = b.getBoundingClientRect();
            const top = document.elementFromPoint(r.left + r.width/2, r.top + r.height/2);
            return {w: r.width, h: r.height, inView: r.top >= 0 && r.bottom <= innerHeight,
                    notCovered: !!top && (b === top || b.contains(top))}; }""")
        check("ปุ่ม 'เริ่มค้นหา' โผล่ ใหญ่ (สูง ≥ 40px) อยู่ในจอ ไม่ถูกบัง",
              vis["h"] >= 40 and vis["inView"] and vis["notCovered"], str(vis))
        time.sleep(2.5)
        check("ยังไม่เข้าสู่ระบบ → ไม่เริ่มค้นเอง", status()["state"] == "awaiting_login")
        S["logged_in"] = True                  # ผู้ใช้เข้าสู่ระบบในหน้าต่าง Chrome
        s = wait_state(lambda s: s["autostart_in"] > 0 or s["state"] == "running", 20, "นับถอยหลัง")
        if s["state"] == "awaiting_login":
            page.wait_for_function(
                "document.querySelector('#btnConfirmLogin').textContent.includes('เริ่มเองใน')",
                timeout=5000)
            check("เข้าสู่ระบบแล้ว → ปุ่มแสดงนับถอยหลัง 'เริ่มเองใน N วิ'", True)
        wait_state(lambda s: s["state"] in ("running", "done"), 20, "เริ่มค้นเอง")
        check("เริ่มค้นเองโดยไม่ต้องกดปุ่ม", True)

        print("\n── ค้นทั้งไฟล์ (มีกล่อง 'มีอะไรใหม่' + ERROR ที่ต้องย้อนกลับไปแก้) ──")
        s = wait_state(lambda s: s["state"] in ("done", "stopped", "error"), 180, "จบรอบ")
        log = "\n".join(s["log"])
        check("รอบจบแบบเสร็จสมบูรณ์", s["state"] == "done", s["message"])
        check("ปิดกล่อง 'มีอะไรใหม่' ของเว็บให้เอง", "ปิดกล่องประกาศ" in log)
        check("แถว ERROR ถูกย้อนกลับไปค้นใหม่ในรอบลองใหม่", "รอบลองใหม่ที่ 1" in log, log[-600:])
        check("ไม่มีแถวผิดพลาดค้าง", s["errors_left"] == 0, str(s["errors_left"]))
        check("แถบความคืบหน้า 100%", s["percent"] == 100, str(s["percent"]))
        hits = dict(S["hits"])
        check("ส่งเลข 12 หลักไปค้นเป็น 13 หลักที่มี 0 ข้างหน้า", hits.get("0123456789012") == 1, str(hits))
        check("ไม่เคยค้นเลขผิด (ศูนย์ต่อท้าย) จากเซลล์ทศนิยม",
              "1234567890130" not in hits and hits.get("0123456789013") == 1, str(hits))
        check("สูตร Excel ใช้ค่าที่คำนวณแล้ว (ไม่ใช่ตัวสูตร)", hits.get("3101501909808") == 1, str(hits))
        check("แถว ERROR! / ERROR#3 จากรอบก่อนถูกค้นใหม่",
              hits.get("3101501909806") == 1 and hits.get("3101501909807") == 1, str(hits))
        check("เว็บตอบ 500 → ย้อนกลับไปค้นใหม่แล้วได้ผล", hits.get("3101501909805") == 2, str(hits))
        check("แถวที่มีผลอยู่แล้วไม่ถูกค้นซ้ำ (ไม่เสียโควตา)", "0987654321098" not in hits, str(hits))
        check("เลขใช้ไม่ได้ไม่ถูกส่งไปค้น", not any(len(k) != 13 for k in hits), str(hits))

        print("\n── ผลในไฟล์ Excel ──")
        import openpyxl
        job = engine.load_queue()[0]
        wb = openpyxl.load_workbook(job["path"])
        ws = wb.active
        B = {r: ws.cell(r, 2).value for r, *_ in ROWS}
        E = {r: ws.cell(r, 5).value for r, *_ in ROWS}
        fill = lambda r: (ws.cell(r, 2).fill.start_color.rgb or "")[-6:]
        check("(4) เลข 12 หลักถูกเติม 0 ข้างหน้า + มาร์คเหลือง (แถวค้น)",
              B[3] == "0123456789012" and fill(3) == "FFFF00", f"{B[3]!r} {fill(3)}")
        check("(4) เลข 12 หลักที่มีผลอยู่แล้วก็ถูกเติม 0 ด้วย",
              B[9] == "0987654321098" and fill(9) == "FFFF00", f"{B[9]!r}")
        check("(4) เลขทศนิยม E+ ถูกเติม 0 ข้างหน้า (ไม่ใช่ต่อท้าย)", B[11] == "0123456789013", repr(B[11]))
        check("สูตรในคอลัมน์เลขบัตรถูกแปลงเป็นค่า (รอบหน้าอ่านได้)", B[10] == "3101501909808", repr(B[10]))
        check("ผลแถว 2: เฉพาะคดีผู้ต้องหา", E[2] == "คดีที่1 ลักทรัพย์ ปี2560", repr(E[2]))
        check("ผลแถว 3 เป็นของคนนี้จริง (ไม่ใช่ผลค้างของแถว 2)", E[3] == "คดีที่1 ยาเสพติด ปี2565", repr(E[3]))
        check("ผลแถว 4 ไม่พบคดี (ไม่ใช่ผลค้างของแถว 3)", E[4] == " - ", repr(E[4]))
        check("(3) แถว 5 (เคย ERROR ในรอบนี้) ได้ผลแล้ว", E[5] == " - ", repr(E[5]))
        check("แถว 6 เลขใช้ไม่ได้ → มีหมายเหตุ", str(E[6]).startswith("⚠ เลขบัตรไม่ถูกต้อง (มี 9 หลัก)"), repr(E[6]))
        check("(3) แถว 7 (ERROR! จากรอบก่อน) ได้ผลแล้ว", E[7] == "คดีที่1 ฉ้อโกง ปี2561", repr(E[7]))
        check("(3) แถว 8 (ERROR#3 จากรอบก่อน) ได้ผลแล้ว", E[8] == "คดีที่1 ทำร้ายร่างกาย ปี", repr(E[8]))
        check("แถว 9 ผลเดิมไม่ถูกแตะ", E[9] == "คดีที่1 เดิม ปี2550", repr(E[9]))
        rowfill = lambda r: [(ws.cell(r, c).fill.start_color.rgb or "")[-6:]
                             if ws.cell(r, c).fill.fill_type else "" for c in range(1, 7)]
        check("v3.5.1 แถว 12 พบคดีแต่ไม่มีข้อหา → ข้อความ (ไม่ระบุข้อหา)",
              E[12] == "คดีที่1 (ไม่ระบุข้อหา) ปี2559", repr(E[12]))
        check("v3.5.1 แถว 12 → เหลืองทั้งแถว A..F รวมเซลล์เลขบัตร",
              rowfill(12) == ["FFE699"] * 6, str(rowfill(12)))
        check("v3.5.1 แถว 8 มีข้อหาแต่ไม่มีปี → แดงอ่อนทั้งแถว ยกเว้นเซลล์เลขบัตร",
              rowfill(8) == ["FFC7CE", "", "FFC7CE", "FFC7CE", "FFC7CE", "FFC7CE"], str(rowfill(8)))
        check("แถวข้อมูลครบไม่ถูกทาสีทั้งแถว", rowfill(2) == [""] * 6, str(rowfill(2)))
        check("สถานะรอบบอกจำนวนแถวไม่มีข้อหา", s["no_charge"] == 1 and s["incomplete"] == 2,
              f"no_charge={s['no_charge']} incomplete={s['incomplete']}")
        check("(3) แถว 11 (ป๊อบอัพค้างครั้งแรก) ได้ผลแล้ว", E[11] == "คดีที่1 ปลอมแปลงเอกสาร ปี2563", repr(E[11]))
        check("ไม่เหลือ ERROR ในไฟล์เลย", not any("ERROR" in str(v or "") for v in E.values()), str(E))
        dl = list(DOWNLOADS.glob("*.xlsx"))
        check("คัดลอกไฟล์ผลไป Downloads อัตโนมัติ", len(dl) == 1, str(dl))

        print("\n── (2) ไฟล์ที่เสร็จแล้วต้องไม่ค้างในคิว ──")
        page.wait_for_selector(".qst-ready", timeout=10000)
        check("หลังจบรอบ ป้ายขึ้น ✓ File ready", "File ready" in page.inner_text("#queueList"))
        page.click('[data-view="dashboard"]')
        page.click('[data-view="run"]')
        page.wait_for_function("document.querySelector('#uploadMsg').textContent.includes('ออกจากคิว')",
                               timeout=8000)
        page.wait_for_function("document.querySelectorAll('#queueList .qitem').length === 0",
                               timeout=8000)
        check("กลับเข้าหน้าอัปโหลด → ไฟล์ที่เสร็จถูกเอาออกจากคิวเอง",
              engine.load_queue() == [] and "ยังไม่มีไฟล์ในคิว" in page.inner_text("#queueList"),
              str(engine.load_queue()))

        print("\n── เริ่มด้วยคิวที่มีแต่ไฟล์ที่เสร็จแล้ว → แจ้งเหตุผลชัดเจน ──")
        engine.save_queue([job])
        r = page.evaluate("""async () => (await (await fetch('/api/run/start', {method:'POST',
            headers:{'Content-Type':'application/json'}, body:'{}'})).json())""")
        check("ไม่เปิดเบราว์เซอร์ซ้ำกับไฟล์ที่เสร็จแล้ว + บอกเหตุผล",
              "ค้นครบแล้ว" in (r.get("error") or ""), str(r))
        engine.save_queue([])

        print("\n── ผู้ใช้ปิดหน้าต่าง Chrome เองระหว่างรอเข้าสู่ระบบ ──")
        S["logged_in"] = False
        src2 = TEST_DATA / "งานสอง.xlsx"
        wb2 = openpyxl.Workbook()
        wb2.active.cell(1, 2, "เลขบัตร")
        wb2.active.cell(1, 5, "ผล")
        for i in range(3):
            wb2.active.cell(2 + i, 2, 3101501909900 + i)
        wb2.save(src2)
        page.set_input_files("#fileInput", str(src2))
        page.wait_for_function("document.querySelector('#uploadMsg').textContent.includes('เพิ่มแล้ว')",
                               timeout=10000)
        # กลับเข้าหน้านี้ใหม่ = ต้องเลือก ②③ ใหม่ทุกครั้ง (กติกา v3.1.0 ไม่ให้ข้ามขั้นตอน)
        page.select_option("#cfgReason", "random")
        page.select_option("#cfgSpeed", "5")
        page.wait_for_function("!document.querySelector('#btnStart').disabled", timeout=5000)
        page.click("#btnStart")
        wait_state(lambda s: s["state"] == "awaiting_login", 60, "awaiting_login (2)")
        time.sleep(1.0)
        kill_crimes_browser()
        s = wait_state(lambda s: s["state"] != "awaiting_login", 20, "ตรวจว่า Chrome ถูกปิด")
        check("ตรวจพบว่า Chrome ถูกปิด → ไม่ค้างสถานะรอ login", s["state"] == "error", s["state"])
        check("บอกให้กด LOGIN CRIMES ใหม่", "LOGIN CRIMES" in s["message"], s["message"])
        page.wait_for_function("!document.querySelector('#btnStart').disabled", timeout=5000)
        check("ปุ่ม LOGIN CRIMES กดได้อีกครั้ง", True)

        print("\n── ผู้ใช้ปิดหน้าต่าง Chrome ระหว่างกำลังค้น ──")
        S["logged_in"] = True
        S["delay"] = 1.5
        page.click("#btnStart")
        wait_state(lambda s: s["state"] == "running" and s["file_done"] >= 1, 60, "ค้นไป 1 แถว")
        kill_crimes_browser()
        s = wait_state(lambda s: s["state"] not in ("running", "awaiting_login", "launching"), 30,
                       "หยุดหลัง Chrome ถูกปิด")
        check("หยุดรอบอย่างสะอาด พร้อมบอกสาเหตุ",
              s["state"] in ("stopped", "error") and "Chrome" in s["message"], s["message"])
        job2 = engine.load_queue()[0]
        ws2 = openpyxl.load_workbook(job2["path"]).active
        res2 = [ws2.cell(2 + i, 5).value for i in range(3)]
        check("ไม่เขียน ERROR ไล่ทุกแถวที่เหลือหลัง Chrome ถูกปิด",
              not any("ERROR" in str(v or "") for v in res2), str(res2))
        check("แถวที่ค้นได้ก่อนปิดถูกบันทึกลงไฟล์แล้ว", res2[0] == " - ", str(res2))

        b.close()

    real = [e for e in JS_ERRORS if "favicon" not in e]
    check("ไม่มี JavaScript error / ไม่มีแจ้งเตือนเริ่มไม่สำเร็จ", not real, "\n    ".join(real[:5]))
    print(f"\nผล: ผ่าน {PASS} · ตก {FAIL}")
    return 1 if FAIL else 0


try:
    code = main()
except AssertionError as e:
    print(f"  ✗ {e}")
    code = 1
finally:
    kill_crimes_browser()
    shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(code)
